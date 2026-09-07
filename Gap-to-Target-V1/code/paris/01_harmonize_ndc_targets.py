# -*- coding: utf-8 -*-
"""
Paris NDC target harmonisation
================

Ce script construit la base maître harmonisée des Nationally Determined
Contributions (NDCs) à partir du fichier source NDCV1.xlsx.

Entrée attendue :
    NDCV1.xlsx

Sorties :
    NDC_Harmonized_Database.xlsx
    NDC_Observed_Emissions_Scope_Harmonized.xlsx

Le classeur de sortie contient trois feuilles :
    1. NDC_Raw_Data   : copie des données source, sans variables analytiques ;
    2. NDC_Harmonized : données enrichies et standardisées pour l'analyse ;
    3. NDC_Codebook   : définition, construction et usage des variables ajoutées.

Principes :
- Les données source sont conservées séparément pour assurer la traçabilité.
- First_NDC et Second_NDC constituent l'échantillon analytique prioritaire,
  sans supprimer les autres générations du fichier harmonisé.
- Les variables standardisées sont ajoutées à droite des colonnes source.
- Le codage est conservateur : une information non suffisamment explicite
  reste manquante / UNKNOWN plutôt que d'être imputée arbitrairement.
- Les corrections manuelles validées lors de la revue pays-par-pays sont
  centralisées dans MANUAL_OVERRIDES.
- La mise en forme du fichier de sortie distingue visuellement données brutes,
  données harmonisées et documentation méthodologique.
"""

from pathlib import Path
import re
import warnings
import unicodedata
import numpy as np
import pandas as pd

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

# 0. FICHIERS

PROJECT_ROOT = Path(__file__).resolve().parents[2]
BASE_DIR = PROJECT_ROOT / "data" / "raw" / "paris"
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed" / "paris"
PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
INPUT_FILE = BASE_DIR / "NDCV1.xlsx"
if not INPUT_FILE.exists():
    alt = BASE_DIR / "NDCV1(2).xlsx"
    if alt.exists():
        INPUT_FILE = alt
OUTPUT_FILE = PROCESSED_DIR / "NDC_Harmonized_Database.xlsx"
INPUT_SHEET = "Feuil1"

# Emissions data used to harmonise the OBSERVED-emissions perimeter
# with the quantified NDC perimeter.
#
# EDGAR:
#   - TOTALS BY COUNTRY = national GHG emissions, Gg CO2eq
#   - IPCC 2006        = sector detail, Gg CO2eq (primary classification)
#   - IPCC 1996        = retained for cross-check / documentation
#
# FAOSTAT:
#   - annual net CO2 emissions/removals from Forest land, Gg CO2
#   - used ONLY as a transparent proxy for the forest/LULUCF component
#     when the NDC explicitly includes net removals/sinks/LULUCF accounting.
EDGAR_FILE = BASE_DIR / "EDGAR_AR5_GHG_1970_2024.xlsx"
FAO_FOREST_FILE = BASE_DIR / "Emissions from forests (Global, National - Annual) - FAOSTAT.xlsx"

EMISSIONS_SCOPE_OUTPUT = PROCESSED_DIR / "NDC_Observed_Emissions_Scope_Harmonized.xlsx"

SHEET_RAW = "NDC_Raw_Data"
SHEET_HARMONIZED = "NDC_Harmonized"
SHEET_CODEBOOK = "NDC_Codebook"

# Palette du classeur
COLOR_RAW = "1E405F"
COLOR_HARMONIZED = "3F6B57"
COLOR_CODEBOOK = "C47A2C"
COLOR_DERIVED_LIGHT = "E7F0EA"
COLOR_CODEBOOK_LIGHT = "F8ECDD"
COLOR_WHITE = "FFFFFF"
COLOR_GRID = "D9E1E8"
COLOR_CHECK = "FFF2CC"
COLOR_LIMITED = "F4CCCC"
COLOR_REVIEWED = "D9EAD3"

# 1. OUTILS GENERAUX

# Canonical emissions unit used throughout the harmonised database and builder.
# 1 MtCO2e = 1,000 GgCO2e; 1 ktCO2e = 1 GgCO2e.
EMISSIONS_UNIT = "GgCO2e"

def mtco2e_to_ggco2e(value):
    """Convert million tonnes CO2e (MtCO2e) to GgCO2e."""
    return np.nan if pd.isna(value) else float(value) * 1000.0

def ktco2e_to_ggco2e(value):
    """Convert thousand tonnes CO2e (ktCO2e) to GgCO2e (numerically identical)."""
    return np.nan if pd.isna(value) else float(value)

def clean_colname(x):
    if x is None:
        return ""
    return str(x).replace("\xa0", " ").strip()

def clean_text(x):
    if pd.isna(x):
        return ""
    x = str(x).replace("\xa0", " ")
    x = re.sub(r"\s+", " ", x).strip()
    return x

def ascii_lower(x):
    s = clean_text(x).lower()
    return "".join(
        c for c in unicodedata.normalize("NFKD", s)
        if not unicodedata.combining(c)
    )

def to_num(x):
    """Conversion prudente des pourcentages / nombres stockés comme texte."""
    if pd.isna(x):
        return np.nan
    if isinstance(x, (int, float, np.number)):
        return float(x)
    s = clean_text(x)
    if not s:
        return np.nan
    s = s.replace("%", "").replace("−", "-").replace("–", "-")
    s = s.replace(",", ".")
    # Conserver uniquement la première valeur numérique si la cellule n'est
    # pas un nombre simple. Les cas complexes sont traités manuellement.
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group()) if m else np.nan

def parse_year(x):
    """Extrait une année 19xx/20xx d'une cellule."""
    if pd.isna(x):
        return np.nan
    if isinstance(x, (int, float, np.number)):
        v = float(x)
        # année directement stockée
        if 1900 <= v <= 2100:
            return int(v)
        # date Excel
        if 20000 <= v <= 70000:
            try:
                return pd.Timestamp("1899-12-30") + pd.to_timedelta(v, unit="D")
            except Exception:
                return np.nan
    s = clean_text(x)
    years = re.findall(r"\b(19\d{2}|20\d{2}|21\d{2})\b", s)
    return int(years[0]) if years else np.nan

def parse_submission_year(x):
    if pd.isna(x):
        return np.nan

    # Date Excel
    if isinstance(x, (int, float, np.number)):
        v = float(x)
        if 1900 <= v <= 2100:
            return int(v)
        if 20000 <= v <= 70000:
            try:
                return int((pd.Timestamp("1899-12-30") +
                            pd.to_timedelta(v, unit="D")).year)
            except Exception:
                pass

    dt = pd.to_datetime(x, errors="coerce")
    if not pd.isna(dt):
        return int(dt.year)

    return parse_year(x)

def parse_timeframe(x):
    """
    Extrait les bornes d'une période d'implémentation.
    La règle utilisée pour la base enrichie est conservatrice :
    - 2 années ou plus : min = début, max = fin
    - une seule année : ne pas inventer une période complète.
    """
    s = clean_text(x)
    if not s:
        return np.nan, np.nan

    years = [int(y) for y in re.findall(r"\b(19\d{2}|20\d{2}|21\d{2})\b", s)]
    if len(years) >= 2:
        return min(years), max(years)
    return np.nan, np.nan

def nonempty(x):
    return clean_text(x) != ""

# 2. CHARGEMENT

# Copie séparée des données source pour la première feuille.
raw_df = pd.read_excel(INPUT_FILE, sheet_name=INPUT_SHEET, engine="openpyxl")

# La feuille harmonisée part d'une copie des données source.
df = raw_df.copy()
df.columns = [clean_colname(c) for c in df.columns]

# Supprimer uniquement dans la feuille harmonisée les colonnes totalement vides.
empty_unnamed = [
    c for c in df.columns
    if (not c or c.lower().startswith("unnamed")) and df[c].isna().all()
]
if empty_unnamed:
    df = df.drop(columns=empty_unnamed)

# Nombre de colonnes source conservées dans NDC_Harmonized.
N_SOURCE_COLUMNS = len(df.columns)

# Récupération robuste des colonnes, avec gestion de l'espace insécable
COL = {clean_colname(c): c for c in df.columns}

def col(name):
    return COL.get(clean_colname(name), clean_colname(name))

# 3. ECHANTILLON FIRST / SECOND NDC

ndc_num = df[col("NDCs_number")].astype("string").str.strip()

df["analysis_ndc12"] = ndc_num.isin(["First_NDC", "Second_NDC"]).astype(int)
df["ndc_generation_num"] = np.select(
    [ndc_num.eq("First_NDC"), ndc_num.eq("Second_NDC")],
    [1, 2],
    default=np.nan
)

# 4. TARGET PRINCIPAL STANDARDISE

main1 = df[col("Target in percent")].map(to_num)
main2 = df[col("Target percentage")].map(to_num)

df["main_target_pct_std"] = main1.combine_first(main2)

if col("main_target_rule") in df.columns:
    df["main_target_rule_std"] = (
        df[col("main_target_rule")]
        .map(clean_text)
        .replace("", np.nan)
        .str.upper()
        .str.replace(r"\s+", "_", regex=True)
    )
else:
    df["main_target_rule_std"] = np.nan

CONSTRUCTED_RULE_PATTERNS = (
    "AVERAGE", "MEAN", "MIDPOINT", "RANGE",
    "SUM_", "ADDITIONAL", "AGGREGAT"
)
rule_text = df["main_target_rule_std"].fillna("")
df["main_target_constructed_flag"] = rule_text.str.contains(
    "|".join(CONSTRUCTED_RULE_PATTERNS), regex=True
).astype(int)

# 5. CONDITIONALITE

cond_source = (
    df[col("Conditionality")].fillna("").map(clean_text)
    + " "
    + (
        df[col("conditionality_structure")].fillna("").map(clean_text)
        if col("conditionality_structure") in df.columns
        else ""
    )
)

def classify_conditionality(x):
    s = ascii_lower(x)

    if not s:
        return "UNKNOWN"

    if (
        "unconditional only" in s
        or "unconditional_only" in s
    ):
        return "UNCONDITIONAL_ONLY"

    if (
        "conditional only" in s
        or "conditional_only" in s
    ):
        return "CONDITIONAL_ONLY"

    if (
        "partially conditional" in s
        or "partially_conditional" in s
    ):
        return "PARTIALLY_CONDITIONAL"

    if (
        "conditional and unconditional" in s
        or "conditional ndc and unconditional ndc" in s
        or "both" in s
        or "domestic_plus_external" in s
        or "unconditional_plus_conditional" in s
        or "unconditional_plus_additional_conditional" in s
        or "conditional_and_unconditional" in s
        or "mixed_conditionality" in s
        or "mixed_sectoral_conditionality" in s
        or "total_with_components" in s
    ):
        return "MIXED_CONDITIONALITY"

    return "UNKNOWN"

df["conditionality_std"] = cond_source.map(classify_conditionality)

df["conditional_support_flag"] = np.select(
    [
        df["conditionality_std"].eq("UNCONDITIONAL_ONLY"),
        df["conditionality_std"].isin(
            ["CONDITIONAL_ONLY", "MIXED_CONDITIONALITY",
             "PARTIALLY_CONDITIONAL"]
        ),
    ],
    [0, 1],
    default=np.nan
)

df["unconditional_pct_std"] = (
    df[col("Unconditional_target in percent")].map(to_num)
)

cond_raw = df[col("Conditional_target in percent")].map(to_num)

# 6. INTERPRETATION DE LA CIBLE CONDITIONNELLE

def conditional_components(row):
    """
    Retourne :
      conditional_total_pct_std
      conditional_additional_pct_std

    La logique dépend de main_target_rule :
    - SUM_UNCOND_AND_ADDITIONAL_COND :
        la colonne conditionnelle est une composante additionnelle ;
        le total = uncond + cond.
    - MAX_WITH_SUPPORT :
        main_target correspond généralement au niveau maximal avec support ;
        si la colonne conditionnelle semble additionnelle, on privilégie le
        main target comme total.
    - DECLARED_TOTAL / TOTAL_WITH_COMPONENTS :
        main_target est le total déclaré.
    - UNCONDITIONAL_PLUS_CONDITIONAL_TOTAL :
        la cible conditionnelle peut déjà être le total.
    - CONDITIONAL_ONLY :
        main target / conditional = cible conditionnelle totale.
    """
    main = row["main_target_pct_std"]
    u = row["unconditional_pct_std"]
    c = row["_cond_raw"]
    rule = clean_text(row["main_target_rule_std"]).upper()

    total = np.nan
    additional = np.nan

    if rule == "SUM_UNCOND_AND_ADDITIONAL_COND":
        if pd.notna(u) and pd.notna(c):
            total = u + c
            additional = c
        elif pd.notna(main):
            total = main
            if pd.notna(u):
                additional = main - u

    elif rule in {"DECLARED_TOTAL", "TOTAL_WITH_COMPONENTS"}:
        if pd.notna(main):
            total = main
            if pd.notna(u):
                additional = main - u
        elif pd.notna(c):
            total = c

    elif rule in {
        "MAX_WITH_SUPPORT",
        "CONDITIONAL_MAX_TARGET",
        "UNCONDITIONAL_PLUS_CONDITIONAL_TOTAL"
    }:
        if pd.notna(main):
            total = main
            if pd.notna(u):
                additional = main - u
        elif pd.notna(c):
            total = c
            if pd.notna(u) and c >= u:
                additional = c - u

    elif rule == "CONDITIONAL_ONLY":
        total = main if pd.notna(main) else c

    elif row["conditionality_std"] == "UNCONDITIONAL_ONLY":
        total = np.nan
        additional = np.nan

    else:
        # Fallback conservateur :
        # si main est une cible totale déclarée et u existe, main est le total.
        if pd.notna(main) and pd.notna(u) and main >= u:
            total = main
            additional = main - u
        elif pd.notna(c):
            total = c
            if pd.notna(u) and c >= u:
                additional = c - u

    if pd.notna(additional) and abs(additional) < 1e-10:
        additional = 0.0

    return pd.Series(
        [total, additional],
        index=["conditional_total_pct_std",
               "conditional_additional_pct_std"]
    )

df["_cond_raw"] = cond_raw
df[["conditional_total_pct_std",
    "conditional_additional_pct_std"]] = df.apply(
        conditional_components, axis=1
    )

df["unconditional_share_main"] = np.where(
    df["main_target_pct_std"].notna()
    & (df["main_target_pct_std"] != 0)
    & df["unconditional_pct_std"].notna(),
    df["unconditional_pct_std"] / df["main_target_pct_std"],
    np.nan
)

# 7. METRIQUE ET REFERENCE DE LA CIBLE

def classify_metric_row(row):
    """
    Classify the quantified target metric using a strict priority hierarchy, including per-capita targets.

    Priority:
      1. explicit harmonised rule (e.g. DECLARED_INTENSITY_TARGET);
      2. Target_Type, which is the most specific target-format field;
      3. GHG_target_type as fallback;
      4. Reference / Cap signals as fallback.

    This prevents intensity targets from being incorrectly coded as
    MULTIPLE_TARGET_TYPES merely because GHG_target_type says "base year".
    """
    rule = clean_text(row.get("main_target_rule_std", "")).upper()
    ttype = ascii_lower(row.get(col("Target_Type"), ""))
    gtype = ascii_lower(row.get(col("GHG_target_type"), ""))
    ref = ascii_lower(row.get(col("Reference"), ""))
    cap = ascii_lower(row.get(col("Cap"), ""))

    if (
        "PER_CAPITA" in rule
        or "PERCAPITA" in rule
        or re.search(r"per\s*-?\s*capita|percapita", ttype)
        or re.search(r"per\s*-?\s*capita|percapita", gtype)
    ):
        return "PER_CAPITA_TARGET"

    if "INTENSITY" in rule or "intensity" in ttype or "intensity" in gtype:
        return "INTENSITY_TARGET"

    if "fixed level" in ttype or "fixed-level" in ttype or "FIXED_LEVEL" in rule:
        return "FIXED_LEVEL_TARGET"

    if (
        "baseline scenario" in ttype
        or "business as usual" in ttype
        or re.search(r"\bbau\b", ttype)
    ):
        return "BAU_REDUCTION"

    if "base year" in ttype:
        return "BASE_YEAR_REDUCTION"

    if "trajectory" in ttype:
        return "TRAJECTORY_TARGET"

    if "no ghg target" in ttype or "not applicable" in ttype:
        return "NO_QUANTIFIED_GHG_TARGET"

    # Fallback to GHG_target_type only when Target_Type was not decisive.
    if "intensity" in gtype:
        return "INTENSITY_TARGET"
    if "fixed level" in gtype or "fixed-level" in gtype:
        return "FIXED_LEVEL_TARGET"
    if (
        "baseline scenario" in gtype
        or "business as usual" in gtype
        or re.search(r"\bbau\b", gtype)
    ):
        return "BAU_REDUCTION"
    if "base year" in gtype:
        return "BASE_YEAR_REDUCTION"
    if "trajectory" in gtype:
        return "TRAJECTORY_TARGET"
    if "no ghg target" in gtype or "not applicable" in gtype:
        return "NO_QUANTIFIED_GHG_TARGET"

    # Last-resort signals.
    if "bau" in ref or "business as usual" in ref:
        return "BAU_REDUCTION"
    if "cap" in cap:
        return "FIXED_LEVEL_TARGET"

    return "OTHER_OR_UNKNOWN"

df["target_metric_std"] = df.apply(classify_metric_row, axis=1)

def target_reference(row):
    metric = row["target_metric_std"]
    ref = ascii_lower(row.get("_reference_raw", ""))

    if metric == "BAU_REDUCTION":
        return "BAU"
    if metric == "BASE_YEAR_REDUCTION":
        return "BASE_YEAR"
    if metric == "INTENSITY_TARGET":
        return "INTENSITY_BASE"
    if metric == "PER_CAPITA_TARGET":
        return "PER_CAPITA_BASE"
    if metric == "FIXED_LEVEL_TARGET":
        return "FIXED_LEVEL"
    if metric == "TRAJECTORY_TARGET":
        return "TRAJECTORY"

    if "bau" in ref or "business as usual" in ref:
        return "BAU"
    if re.search(r"\b(19|20)\d{2}\b", ref):
        return "BASE_YEAR"
    return "UNKNOWN"

df["_reference_raw"] = df[col("Reference")].fillna("")
df["target_reference_std"] = df.apply(target_reference, axis=1)

# 8. PORTEE ET SECTEURS

sector_text = df[col("Sectors")].fillna("").map(clean_text)
sector_text_ascii = sector_text.map(ascii_lower)

def has_any(s, patterns):
    return int(any(re.search(p, s) for p in patterns))

sector_patterns = {
    "sector_energy_flag": [
        r"\benergy\b", r"\bpower\b", r"electric", r"fuel combustion"
    ],
    "sector_industry_ipu_flag": [
        r"\bindustry\b", r"industrial", r"\bippu\b",
        r"industrial process", r"\bpiup\b"
    ],
    "sector_agriculture_afolu_flag": [
        r"\bagriculture\b", r"\bafolu\b", r"\bafat\b",
        r"livestock", r"cropland"
    ],
    "sector_lulucf_forestry_flag": [
        r"\blulucf\b", r"\blucf\b", r"forestr", r"land use",
        r"forest land"
    ],
    "sector_waste_flag": [
        r"\bwaste\b", r"wastewater", r"landfill"
    ],
    "sector_transport_flag": [
        r"\btransport\b", r"transportation", r"mobility"
    ],
}

for var, patterns in sector_patterns.items():
    df[var] = sector_text_ascii.map(lambda s: has_any(s, patterns))

sector_flag_cols = list(sector_patterns.keys())
df["sector_count_major"] = df[sector_flag_cols].sum(axis=1)

def classify_scope(row):
    """
    Descriptive coverage based on the sectors field.
    This variable is NOT by itself an exclusion rule.
    """
    s = ascii_lower(row.get("_sector_text", ""))
    n = row["sector_count_major"]
    rule = clean_text(row.get("main_target_rule_std", "")).upper()

    if any(k in s for k in [
        "economy-wide", "economy wide", "entire economy",
        "all sectors", "all ipcc sectors", "whole economy",
        "all categories", "all emission sectors",
        "100% of national emissions", "100 percent of national emissions",
        "all the economic sectors", "nationwide"
    ]):
        return "ECONOMY_WIDE_OR_BROAD"

    if "SECTORAL_TARGET" in rule:
        return "SECTORAL"

    if n >= 4:
        return "ECONOMY_WIDE_OR_BROAD"
    if 2 <= n <= 3:
        return "MULTISECTORAL"
    if n == 1:
        return "SECTORAL"
    return "UNKNOWN"

df["_sector_text"] = sector_text
df["target_scope_std"] = df.apply(classify_scope, axis=1)

def classify_quantified_target_coverage(row):
    """
    Coverage of the QUANTIFIED TARGET itself, not merely the sectors where
    mitigation measures are implemented.

    This distinction is essential for comparability with national EDGAR GHG
    emissions.
    """
    s = ascii_lower(row.get("_sector_text", ""))
    rule = clean_text(row.get("main_target_rule_std", "")).upper()
    n = row["sector_count_major"]

    # Explicit national/economy-wide wording has highest priority.
    if any(k in s for k in [
        "economy-wide", "economy wide", "entire economy",
        "all sectors", "all ipcc sectors", "all emission sectors",
        "whole economy", "all categories",
        "100% of national emissions", "100 percent of national emissions",
        "all the economic sectors", "nationwide",
        "all emitting sectors"
    ]):
        return "NATIONAL_GHG"

    # Explicitly constructed sectoral targets cannot be treated as national.
    if "MEAN_SECTORAL" in rule or "AVERAGE_SECTORAL" in rule:
        return "PARTIAL_SECTORAL_GHG"
    if "SECTORAL_TARGET" in rule:
        return "SINGLE_SECTOR_GHG"

    # Strong single-sector wording.
    single_sector_phrases = [
        "electricity",
        "electricity generation",
        "energy sector commitment",
        "transportation (for unconditional target)",
        "transport sector",
        "power sector"
    ]
    if n == 1 and any(k in s for k in single_sector_phrases):
        return "SINGLE_SECTOR_GHG"

    # Broad inventory-like coverage even without the literal phrase
    # "economy-wide".
    if n >= 4:
        return "BROAD_NATIONAL_GHG"

    # Several sectors may still represent a national target, so do not
    # automatically exclude them.
    if 2 <= n <= 3:
        return "UNCLEAR"

    if n == 1:
        return "UNCLEAR"

    return "UNCLEAR"

df["quantified_target_coverage_std"] = df.apply(
    classify_quantified_target_coverage, axis=1
)

# 9. FLAGS SPECIAUX

combined_text = (
    df[col("Additional infomation")].fillna("").map(clean_text)
    + " "
    + df[col("Sectors")].fillna("").map(clean_text)
    + " "
    + df[col("Reference")].fillna("").map(clean_text)
)

df["per_capita_target_flag"] = (
    combined_text.map(ascii_lower)
    .str.contains(r"per capita|per-capita|percapita", regex=True, na=False)
    | df["target_metric_std"].eq("PER_CAPITA_TARGET")
).astype(int)

df["net_emissions_target_flag"] = (
    combined_text.map(ascii_lower)
    .str.contains(r"\bnet emissions\b|\bnet ghg\b|net greenhouse", regex=True, na=False)
).astype(int)

df["target_range_flag"] = (
    df["main_target_rule_std"].fillna("")
    .str.contains("RANGE|AVERAGE|MEAN|MIDPOINT", regex=True)
).astype(int)

# 10. PERIODE D'IMPLEMENTATION

timeframe = df[col("Timeframe for implementation")]
parsed = timeframe.map(parse_timeframe)
df["implementation_start_year"] = [x[0] for x in parsed]
df["implementation_end_year"] = [x[1] for x in parsed]

df["implementation_duration_years"] = np.where(
    df["implementation_start_year"].notna()
    & df["implementation_end_year"].notna(),
    df["implementation_end_year"]
    - df["implementation_start_year"]
    + 1,
    np.nan
)

interim_pct_col = col("interim_target in percent")
interim_year_col = col("interim_target_year")

df["interim_target_flag"] = (
    df[interim_pct_col].map(to_num).notna()
    | df[interim_year_col].map(parse_year).notna()
).astype(int)

# 11. CIBLES ABSOLUES / ANNEES

df["main_target_value_ggco2e_std"] = (
    df[col("Target_in_value_GgCO₂eq")].map(to_num)
)
df["unconditional_target_value_ggco2e_std"] = (
    df[col("Unconditional_target in value_GgCO₂eq")].map(to_num)
)
df["conditional_target_value_ggco2e_std"] = (
    df[col("Conditional_target in value_GgCO₂eq")].map(to_num)
)

def absolute_kind(row):
    if pd.isna(row["main_target_value_ggco2e_std"]):
        return np.nan

    metric = row["target_metric_std"]
    cap = ascii_lower(row.get("_cap_raw", ""))
    rule = clean_text(row.get("main_target_rule_std", "")).upper()
    info = ascii_lower(row.get(col("Additional infomation"), ""))

    # Carbon-budget/cumulative amounts must never be treated as an end-year cap.
    if any(k in info for k in [
        "carbon budget", "cumulative emissions", "cumulative emission",
        "emissions budget", "emission budget"
    ]):
        return "CUMULATIVE_EMISSIONS_BUDGET"

    if metric == "FIXED_LEVEL_TARGET" or "cap" in cap or "FIXED_LEVEL" in rule:
        return "EMISSIONS_LEVEL_OR_CAP"

    # A BAU-referenced value with no percentage is interpreted as an
    # absolute reduction amount only when the target architecture says so.
    if (
        "ABSOLUTE" in rule
        or "reduction" in cap
        or (
            metric == "BAU_REDUCTION"
            and pd.isna(row.get("main_target_pct_std", np.nan))
        )
    ):
        return "ABSOLUTE_REDUCTION_AMOUNT"

    return "REPORTED_ABSOLUTE_TARGET_OR_REDUCTION"

df["_cap_raw"] = df[col("Cap")].fillna("")
df["main_target_value_kind_std"] = df.apply(absolute_kind, axis=1)

df["absolute_target_available_flag"] = (
    df["main_target_value_ggco2e_std"].notna().astype(int)
)

df["base_year_std"] = df[col("Base_year")].map(parse_year)
df["target_year_std"] = df[col("Target_year")].map(parse_year)

# Use an explicit object/string Series instead of np.where mixing
# strings and np.nan. Recent NumPy versions raise a DTypePromotionError
# when attempting to combine StrDType and float NaN in np.where.
df["target_year_source"] = pd.Series(
    pd.NA,
    index=df.index,
    dtype="object"
)
df.loc[
    df["target_year_std"].notna(),
    "target_year_source"
] = "DECLARED_TARGET_YEAR"

# Recover a missing target year from the explicit implementation timeframe
# only when the end year is available and unambiguous.
recover_target_year = (
    df["target_year_std"].isna()
    & df["implementation_end_year"].notna()
    & df["implementation_start_year"].notna()
)
df.loc[recover_target_year, "target_year_std"] = (
    df.loc[recover_target_year, "implementation_end_year"]
)
df.loc[recover_target_year, "target_year_source"] = "TIMEFRAME_END_RECOVERED"

df["submission_year_std"] = df[col("Submission-date")].map(parse_submission_year)

df["target_horizon_years"] = np.where(
    df["target_year_std"].notna() & df["submission_year_std"].notna(),
    df["target_year_std"] - df["submission_year_std"],
    np.nan
)

df["post2030_target_flag"] = (
    (df["target_year_std"] > 2030).fillna(False).astype(int)
)

# 12. TARGET QUANTIFIE ET ECHANTILLON STRICT

df["quantified_ghg_target_flag"] = (
    df["main_target_pct_std"].notna()
    | df["main_target_value_ggco2e_std"].notna()
    | df["target_metric_std"].isin(
        ["FIXED_LEVEL_TARGET", "TRAJECTORY_TARGET"]
    )
).astype(int)

df["strict_comparable_pct_target_flag"] = (
    df["analysis_ndc12"].eq(1)
    & df["main_target_pct_std"].notna()
    & df["target_metric_std"].isin(
        ["BAU_REDUCTION", "BASE_YEAR_REDUCTION"]
    )
    & df["target_scope_std"].eq("ECONOMY_WIDE_OR_BROAD")
    & df["per_capita_target_flag"].eq(0)
).astype(int)

# 13. CORRECTIONS MANUELLES VALIDEES PAYS-PAR-PAYS
#
# Ces corrections correspondent aux cas explicitement relus et validés
# pendant la revue manuelle des NDC. Les valeurs non listées restent issues
# des colonnes originales et des règles génériques ci-dessus.
#
# Le dictionnaire peut être prolongé facilement.

MANUAL_OVERRIDES = {
    # pays, cycle : champs standardisés à écraser
    ("Vietnam", "First_NDC"): {
        "implementation_start_year": 2021,
        "implementation_end_year": 2030,
        "target_year_std": 2030,
        "target_year_source": "MANUAL_VERIFIED",
        "quantified_target_coverage_std": "NATIONAL_GHG",
    },
    ("Viet Nam", "First_NDC"): {
        "implementation_start_year": 2021,
        "implementation_end_year": 2030,
        "target_year_std": 2030,
        "target_year_source": "MANUAL_VERIFIED",
        "quantified_target_coverage_std": "NATIONAL_GHG",
    },
    ("Vietnam", "Second_NDC"): {
        "main_target_pct_std": 43.5,
        "implementation_start_year": 2021,
        "implementation_end_year": 2030,
        "target_year_std": 2030,
        "target_year_source": "MANUAL_VERIFIED",
        "quantified_target_coverage_std": "NATIONAL_GHG",
    },
    ("Viet Nam", "Second_NDC"): {
        "main_target_pct_std": 43.5,
        "implementation_start_year": 2021,
        "implementation_end_year": 2030,
        "target_year_std": 2030,
        "target_year_source": "MANUAL_VERIFIED",
        "quantified_target_coverage_std": "NATIONAL_GHG",
    },
    ("Colombia", "First_NDC"): {
        "quantified_target_coverage_std": "NATIONAL_GHG",
    },
    ("Colombia", "Second_NDC"): {
        "quantified_target_coverage_std": "NATIONAL_GHG",
    },
    ("Costa Rica", "First_NDC"): {
        "quantified_target_coverage_std": "NATIONAL_GHG",
    },
    ("Costa Rica", "Second_NDC"): {
        "quantified_target_coverage_std": "BROAD_NATIONAL_GHG",
    },
    ("Saint Kitts and Nevis", "First_NDC"): {
        "quantified_target_coverage_std": "NATIONAL_GHG",
    },
    ("Saint Kitts and Nevis", "Second_NDC"): {
        "quantified_target_coverage_std": "NATIONAL_GHG",
    },
    ("Togo", "Second_NDC"): {
        "main_target_pct_std": 50.57,
    },
    ("Solomon Islands", "First_NDC"): {
        "main_target_pct_std": 75.0,
    },
    ("Solomon Islands", "Second_NDC"): {
        "main_target_pct_std": 78.0,
    },
    ("Vanuatu", "First_NDC"): {
        "main_target_pct_std": 100.0,
        "main_target_rule_std": "SECTORAL_TARGET",
        "target_scope_std": "SECTORAL",
    },
    ("Tuvalu", "First_NDC"): {
        "main_target_pct_std": 60.0,
        "target_scope_std": "SECTORAL",
        "target_year_std": 2025,
        "implementation_start_year": 2020,
        "implementation_end_year": 2025,
    },
    ("Mexico", "Second_NDC"): {
        "main_target_pct_std": 40.0,
    },
}

country_col = col("Country")

reviewed_keys = set()
for (country, cycle), updates in MANUAL_OVERRIDES.items():
    mask = (
        df[country_col].map(clean_text).eq(country)
        & df[col("NDCs_number")].map(clean_text).eq(cycle)
    )
    if mask.any():
        reviewed_keys.add((country, cycle))
        for k, v in updates.items():
            df.loc[mask, k] = v

# Recalculs dépendants après overrides
df["main_target_constructed_flag"] = (
    df["main_target_rule_std"].fillna("")
    .str.contains(
        "AVERAGE|MEAN|MIDPOINT|RANGE|SUM_|ADDITIONAL|AGGREGAT",
        regex=True
    )
    .astype(int)
)

df["target_range_flag"] = (
    df["main_target_rule_std"].fillna("")
    .str.contains("RANGE|AVERAGE|MEAN|MIDPOINT", regex=True)
    .astype(int)
)

df["implementation_duration_years"] = np.where(
    df["implementation_start_year"].notna()
    & df["implementation_end_year"].notna(),
    df["implementation_end_year"]
    - df["implementation_start_year"]
    + 1,
    np.nan
)

df["post2030_target_flag"] = (
    (df["target_year_std"] > 2030).fillna(False).astype(int)
)

# 13B. USABILITE DE LA CIBLE POUR LE GAP NATIONAL

def target_usability_status(row):
    if row.get("analysis_ndc12", 0) != 1:
        return "OUTSIDE_NDC12"

    metric = clean_text(row.get("target_metric_std", ""))
    coverage = clean_text(row.get("quantified_target_coverage_std", ""))
    T = row.get("target_year_std", np.nan)
    s = row.get("implementation_start_year", np.nan)
    pct = row.get("main_target_pct_std", np.nan)
    aval = row.get("main_target_value_ggco2e_std", np.nan)
    akind = clean_text(row.get("main_target_value_kind_std", ""))

    if pd.notna(s) and s > 2024:
        return "NOT_YET_OBSERVABLE"

    if coverage in {"SINGLE_SECTOR_GHG", "PARTIAL_SECTORAL_GHG"}:
        return "SECTOR_SCOPE_INCOMPATIBLE"

    if metric == "INTENSITY_TARGET":
        return "USABLE_AFTER_INTENSITY_CONVERSION"

    if metric == "PER_CAPITA_TARGET":
        return "USABLE_AFTER_PER_CAPITA_CONVERSION"

    if akind == "CUMULATIVE_EMISSIONS_BUDGET":
        return "MANUAL_REVIEW_CUMULATIVE_BUDGET"

    if akind == "ABSOLUTE_REDUCTION_AMOUNT":
        if pd.notna(T):
            return "USABLE_AFTER_ABSOLUTE_REDUCTION_CONVERSION"
        return "MISSING_TARGET_YEAR"

    if akind == "EMISSIONS_LEVEL_OR_CAP":
        if pd.notna(T):
            return "USABLE_FIXED_LEVEL"
        return "MISSING_TARGET_YEAR"

    if metric in {"BAU_REDUCTION", "BASE_YEAR_REDUCTION"}:
        if pd.isna(T):
            return "MISSING_TARGET_YEAR"
        if pd.notna(pct):
            if coverage == "UNCLEAR":
                return "USABLE_SCOPE_REVIEW"
            return "USABLE"

    if metric == "NO_QUANTIFIED_GHG_TARGET":
        return "NO_QUANTIFIED_GHG_TARGET"

    if pd.isna(pct) and pd.isna(aval):
        return "INSUFFICIENT_INFORMATION"

    return "MANUAL_REVIEW"

df["target_usability_status"] = df.apply(target_usability_status, axis=1)

# DOCUMENTED COUNTRY-SPECIFIC TARGET CLARIFICATIONS
# These overrides are deliberately narrow and auditable. They correct
# ambiguities in structured source fields without changing the general
# classification rules.

def _country_cycle_mask(country_pattern, cycle_pattern):
    """
    Robust country-cycle selector.

    The raw NDC database uses the column `NDCs_number`
    (e.g. First_NDC, Second_NDC), not `NDC`.
    """
    country_txt = df[col("Country")].astype(str).str.strip()
    cycle_txt = df[col("NDCs_number")].astype(str).str.strip()

    return (
        country_txt.str.contains(
            country_pattern, case=False, regex=True, na=False
        )
        & cycle_txt.str.contains(
            cycle_pattern, case=False, regex=True, na=False
        )
    )

# Zimbabwe — updated/revised NDC (2021):
# 40% reduction in GHG emissions per capita relative to BAU by 2030.
# Coverage: Energy, Waste, IPPU and AFOLU; treat as broad/economy-wide
# quantified GHG coverage for the national Gap exercise.
m_zwe_n2 = _country_cycle_mask(
    r"^Zimbabwe$",
    r"^Second_NDC$"
)
if m_zwe_n2.any():
    df.loc[m_zwe_n2, "target_metric_std"] = "BAU_PER_CAPITA_TARGET"
    df.loc[m_zwe_n2, "target_reference_std"] = "BAU_PER_CAPITA"
    df.loc[m_zwe_n2, "quantified_target_coverage_std"] = "ECONOMY_WIDE_OR_BROAD"
    df.loc[m_zwe_n2, "target_usability_status"] = "USABLE"

# Zimbabwe — first NDC:
# quantified mitigation target is energy-sector specific. Preserve exclusion
# from the national economy-wide Gap.
m_zwe_n1 = _country_cycle_mask(
    r"^Zimbabwe$",
    r"^First_NDC$"
)
if m_zwe_n1.any():
    df.loc[m_zwe_n1, "quantified_target_coverage_std"] = "SINGLE_SECTOR_GHG"
    df.loc[m_zwe_n1, "target_usability_status"] = "EXCLUDE_SECTOR_SCOPE"

# Chile — updated NDC (2020):
# The NDC contains BOTH a cumulative 2020–2030 budget of 1,100 MtCO2e and
# a single-year 2030 emissions level of 95 MtCO2e. The Gap trajectory must
# use the single-year terminal level, not the cumulative budget.
m_chl_n2 = _country_cycle_mask(
    r"^Chile$",
    r"^Second_NDC$"
)
if m_chl_n2.any():
    df.loc[m_chl_n2, "target_metric_std"] = "FIXED_LEVEL_TARGET"
    df.loc[m_chl_n2, "target_reference_std"] = "ABSOLUTE_LEVEL"
    df.loc[m_chl_n2, "main_target_value_kind_std"] = "EMISSIONS_LEVEL_OR_CAP"
    df.loc[m_chl_n2, "main_target_value_ggco2e_std"] = 95000.0
    df.loc[m_chl_n2, "target_year_std"] = 2030.0
    df.loc[m_chl_n2, "target_year_source"] = "DOCUMENTED_NDC_LEVEL"
    df.loc[m_chl_n2, "quantified_target_coverage_std"] = "ECONOMY_WIDE_OR_BROAD"
    df.loc[m_chl_n2, "target_usability_status"] = "USABLE"

# FINAL AUDIT OVERRIDES — COUNTRY-BY-COUNTRY VALIDATION
#
# These variables preserve the distinction between:
#   * what the NDC explicitly declares;
#   * how the target is harmonised for the Gap-to-Target;
#   * whether the cycle is included in the main construction.
#
# IMPORTANT:
# Sectoral targets may be INCLUDED when an explicit harmonisation rule was
# validated during the audit. Their actual sector coverage remains stored in
# quantified_target_coverage_std for later cluster / heterogeneity analysis.

# Create text and numeric audit variables with explicit dtypes.
# Pandas 2.x/3.x no longer permits assigning strings into float64 columns
# without an explicit dtype conversion.
for _name in [
    "audit_decision_std",
    "audit_reason_std",
    "audit_note_std",
    "harmonization_method_std",
]:
    if _name not in df.columns:
        df[_name] = pd.Series(pd.NA, index=df.index, dtype="object")
    elif not (
        df[_name].dtype == "object"
        or pd.api.types.is_string_dtype(df[_name])
    ):
        df[_name] = df[_name].astype("object")

for _name in [
    "official_bau_value_ggco2e_std",
    "special_target_amount_ggco2e_std",
    "special_unconditional_target_amount_ggco2e_std",
    "special_conditional_target_amount_ggco2e_std",
    "interim_target_year_std",
    "interim_target_pct_std",
    "interim_unconditional_pct_std",
    "interim_conditional_total_pct_std",
]:
    if _name not in df.columns:
        df[_name] = np.nan

def apply_final_audit_override(country_pattern, cycle, **updates):
    """
    Apply one transparent country/cycle audit decision.

    The helper is dtype-safe with recent pandas versions:
    text overrides force the target column to object/string-compatible dtype
    before assignment, while numeric overrides remain numeric.
    """
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="This pattern is interpreted as a regular expression, and has match groups.*",
            category=UserWarning,
        )
        _country_match = df[col("Country")].astype(str).str.strip().str.contains(
            country_pattern, case=False, regex=True, na=False
        )

    mask = (
        _country_match
        & df[col("NDCs_number")].astype(str).str.strip().eq(cycle)
    )

    if mask.any():
        for key, value in updates.items():

            # Create missing columns with a dtype compatible with the value.
            if key not in df.columns:
                if isinstance(value, str):
                    df[key] = pd.Series(pd.NA, index=df.index, dtype="object")
                else:
                    df[key] = np.nan

            # Recent pandas raises TypeError instead of silently upcasting
            # float columns when a string is assigned.
            if isinstance(value, str):
                if not (
                    df[key].dtype == "object"
                    or pd.api.types.is_string_dtype(df[key])
                ):
                    df[key] = df[key].astype("object")

            df.loc[mask, key] = value

    return int(mask.sum())

# Previously validated fixed/special targets

# Armenia First NDC:
# 633 MtCO2e is a cumulative 2015-2050 emissions budget.
# Validated harmonisation: annualise uniformly over the 36-year budget horizon
# and use the 2015-2030 implementation window for the annual Gap trajectory.
ARM_BUDGET_ANNUAL_GG = mtco2e_to_ggco2e(633.0) / (2050 - 2015 + 1)
apply_final_audit_override(
    r"^Armenia$", "First_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="ANNUALISED_CUMULATIVE_EMISSIONS_BUDGET",
    audit_note_std=(
        "633 MtCO2e is a cumulative 2015-2050 emissions budget; "
        "converted to a uniform annual equivalent for the 2015-2030 Gap window."
    ),
    harmonization_method_std="ANNUALISED_EMISSIONS_BUDGET",
    target_metric_std="FIXED_LEVEL_TARGET",
    target_reference_std="FIXED_LEVEL",
    main_target_value_kind_std="EMISSIONS_LEVEL_OR_CAP",
    main_target_value_ggco2e_std=ARM_BUDGET_ANNUAL_GG,
    implementation_start_year=2015,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)

# Armenia updated/Second NDC: 40% below 1990 by 2030.
apply_final_audit_override(
    r"^Armenia$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_BASE_YEAR_TARGET",
    harmonization_method_std="DIRECT_BASE_YEAR_REDUCTION",
    target_metric_std="BASE_YEAR_REDUCTION",
    target_reference_std="BASE_YEAR",
    main_target_pct_std=40.0,
    base_year_std=1990,
    implementation_start_year=2021,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)

# Guyana: no national quantified GHG target; 52 Mt is a mitigation contribution.
apply_final_audit_override(
    r"^Guyana$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    audit_note_std="52 MtCO2e refers to mitigation contribution, not terminal emissions.",
    harmonization_method_std="NONE",
)

# Belize First NDC: action-based contribution without an aggregate GHG target.
apply_final_audit_override(
    r"^Belize$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)

# Belize Second NDC: use the reported terminal-year annual abatement, not the
# cumulative 2021-2030 amount. The NDC is conditional on external support.
apply_final_audit_override(
    r"^Belize$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="ABSOLUTE_ANNUAL_ABATEMENT_TARGET",
    audit_note_std=(
        "2030 target = scope-consistent BAU emissions minus 1,080 GgCO2e. "
        "The reported 5,647 GgCO2e is cumulative over 2021-2030 and is not "
        "used as a terminal-year subtraction."
    ),
    harmonization_method_std="ABSOLUTE_ANNUAL_ABATEMENT_BAU",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    special_target_amount_ggco2e_std=1080.0,
    special_conditional_target_amount_ggco2e_std=1080.0,
    implementation_start_year=2021,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
)

# Gabon First NDC: 50% below baseline in 2025.
apply_final_audit_override(
    r"^Gabon$", "First_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_BAU_TARGET",
    harmonization_method_std="DIRECT_BAU_REDUCTION",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=50.0,
    unconditional_pct_std=50.0,
    implementation_start_year=2010,
    implementation_end_year=2025,
    target_year_std=2025,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
)
# Gabon Second NDC: net-zero / net-removal target not comparable to gross series.
apply_final_audit_override(
    r"^Gabon$", "Second_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NET_ZERO_TARGET_NOT_COMPARABLE_WITH_GROSS_EMISSIONS_SERIES",
    audit_note_std=(
        "2050 carbon neutrality; conditional 100 MtCO2/year is net uptake, "
        "not a positive fixed emissions level."
    ),
    harmonization_method_std="NONE",
)

# Peru First NDC: 20% unconditional + 10 percentage points conditional = 30% total.
apply_final_audit_override(
    r"^Peru$", "First_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_BAU_TARGET",
    harmonization_method_std="DIRECT_BAU_REDUCTION",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=30.0,
    unconditional_pct_std=20.0,
    conditional_total_pct_std=30.0,
    conditional_additional_pct_std=10.0,
    implementation_start_year=2021,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)

# Peru updated/Second NDC: 208.8 Mt unconditional, 179 Mt conditional/maximum.
apply_final_audit_override(
    r"^Peru$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_FIXED_LEVEL_TARGET",
    harmonization_method_std="DIRECT_ABSOLUTE_LEVEL",
    target_metric_std="FIXED_LEVEL_TARGET",
    target_reference_std="FIXED_LEVEL",
    main_target_value_kind_std="EMISSIONS_LEVEL_OR_CAP",
    main_target_value_ggco2e_std=mtco2e_to_ggco2e(179.0),
    unconditional_target_value_ggco2e_std=mtco2e_to_ggco2e(208.8),
    conditional_target_value_ggco2e_std=mtco2e_to_ggco2e(179.0),
    implementation_start_year=2021,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)

# Bhutan: compare gross EDGAR emissions with the quantified neutrality ceiling
# of 6.3 MtCO2e. FAOSTAT Forest Land is not added because the forest sink is
# already the quantity used to define this emissions cap.
for _cycle, _start in [("First_NDC", 2020), ("Second_NDC", 2021)]:
    apply_final_audit_override(
        r"^Bhutan$", _cycle,
        audit_decision_std="INCLUDE",
        audit_reason_std="VERIFIED_CARBON_NEUTRALITY_GROSS_EMISSIONS_CAP",
        audit_note_std=(
            "Gross GHG emissions must not exceed documented forest "
            "sequestration of 6.3 MtCO2e; use a 6,300 GgCO2e fixed cap "
            "without adding FAOSTAT Forest Land to observed emissions."
        ),
        harmonization_method_std="DIRECT_ABSOLUTE_LEVEL",
        target_metric_std="FIXED_LEVEL_TARGET",
        target_reference_std="FIXED_LEVEL",
        main_target_value_kind_std="EMISSIONS_LEVEL_OR_CAP",
        main_target_value_ggco2e_std=6300.0,
        implementation_start_year=_start,
        implementation_end_year=2030,
        target_year_std=2030,
        target_year_source="MANUAL_AUDIT_VERIFIED",
        quantified_target_coverage_std="BROAD_NATIONAL_GHG",
    )

# Full audit list

# Antigua and Barbuda
for _cycle in ["First_NDC", "Second_NDC"]:
    apply_final_audit_override(
        r"^Antigua (and|&) Barbuda$", _cycle,
        audit_decision_std="EXCLUDE",
        audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
        harmonization_method_std="NONE",
    )

# Bahrain
for _cycle in ["First_NDC", "Second_NDC"]:
    apply_final_audit_override(
        r"^Bahrain$", _cycle,
        audit_decision_std="EXCLUDE",
        audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
        harmonization_method_std="NONE",
    )

# Bolivia First NDC
apply_final_audit_override(
    r"^Bolivia$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)

# Cook Islands First NDC: quantified electricity-sector target.
apply_final_audit_override(
    r"^Cook Islands$", "First_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_SECTOR_MATCHED_TARGET",
    audit_note_std=(
        "Electricity target: 38 percentage points unconditional plus "
        "43 additional conditional = 81% total by 2030; observed emissions "
        "must use EDGAR POWER_HEAT only."
    ),
    harmonization_method_std="SUM_UNCOND_AND_ADDITIONAL_COND",
    main_target_rule_std="SUM_UNCOND_AND_ADDITIONAL_COND",
    target_metric_std="BASE_YEAR_REDUCTION",
    target_reference_std="BASE_YEAR",
    main_target_pct_std=81.0,
    unconditional_pct_std=38.0,
    conditional_total_pct_std=81.0,
    conditional_additional_pct_std=43.0,
    base_year_std=2006,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="SINGLE_SECTOR_GHG",
    target_usability_status="USABLE_SECTOR_MATCHED",
)

# Federated States of Micronesia — the First NDC target is quantified, but
# EDGAR does not report the electricity/heat and transport components needed
# to reproduce its sectoral perimeter.  This is a data-availability exclusion,
# not a "no quantified target" exclusion.
apply_final_audit_override(
    r"^Micronesia$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="MISSING_SCOPE_COMPATIBLE_EDGAR_DATA",
    audit_note_std=(
        "28% unconditional plus 7 additional conditional = 35% total "
        "below 2000 for electricity generation and transport by 2025. "
        "The target is quantified, but EDGAR contains no usable POWER_HEAT "
        "or TRANSPORT series for FSM, so scope-consistent observed emissions "
        "and a Gap-to-Target indicator cannot be constructed."
    ),
    harmonization_method_std="SUM_UNCOND_AND_ADDITIONAL_COND",
    main_target_rule_std="SUM_UNCOND_AND_ADDITIONAL_COND",
    target_metric_std="BASE_YEAR_REDUCTION",
    target_reference_std="BASE_YEAR",
    main_target_pct_std=35.0,
    unconditional_pct_std=28.0,
    conditional_total_pct_std=35.0,
    conditional_additional_pct_std=7.0,
    base_year_std=2000,
    implementation_start_year=2020,
    implementation_end_year=2025,
    target_year_std=2025,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
    target_usability_status="MISSING_SCOPE_COMPATIBLE_EDGAR_DATA",
)
apply_final_audit_override(
    r"^Micronesia$", "Second_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_AGGREGATE_GHG_TARGET",
    audit_note_std=(
        "The supplied updated-NDC record identifies a base-year architecture "
        "but provides no aggregate reduction rate, base year, or fixed "
        "emissions level from which E_target can be reproduced."
    ),
    harmonization_method_std="NONE",
    target_metric_std="NO_QUANTIFIED_GHG_TARGET",
    target_reference_std="UNKNOWN",
    main_target_pct_std=np.nan,
    unconditional_pct_std=np.nan,
    conditional_total_pct_std=np.nan,
    base_year_std=np.nan,
    implementation_start_year=2021,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="UNQUANTIFIED_PARTIAL_SECTORS",
    target_usability_status="NO_QUANTIFIED_TARGET",
)

# Bosnia and Herzegovina — the NDC1 target is numerically explicit but cannot
# be reconstructed under the adopted land-use perimeter because FAOSTAT Forest
# Land is unavailable for its 1990 base year.  NDC2 explicitly excludes sinks
# and is reproducible from the 2014 gross EDGAR national total.
apply_final_audit_override(
    r"^Bosnia and Herzegovina$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="MISSING_FAOSTAT_FOREST_BASE_YEAR",
    audit_note_std=(
        "Use the NDC's explicit base-year equivalents: unconditional "
        "emissions are 18% above 1990 and conditional emissions are 3% "
        "below 1990. The NDC1 accounting scope includes LULUCF sinks, but "
        "FAOSTAT Forest Land is unavailable for 1990; the scope-consistent "
        "base-year emissions level therefore cannot be constructed."
    ),
    harmonization_method_std="DIRECT_BASE_YEAR_REDUCTION",
    target_metric_std="BASE_YEAR_REDUCTION",
    target_reference_std="BASE_YEAR",
    main_target_pct_std=3.0,
    unconditional_pct_std=-18.0,
    conditional_total_pct_std=3.0,
    base_year_std=1990,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
    target_usability_status="MISSING_FAOSTAT_FOREST_BASE_YEAR",
)
apply_final_audit_override(
    r"^Bosnia and Herzegovina$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_BASE_YEAR_TARGET_EXCLUDING_SINKS",
    audit_note_std=(
        "Headline 2030 targets are 12.8% unconditional and 17.5% "
        "conditional below 2014, explicitly excluding GHG sinks. The "
        "separate 93 GgCO2e forestry-sink action is not folded into the "
        "principal target level."
    ),
    harmonization_method_std="DIRECT_BASE_YEAR_REDUCTION",
    target_metric_std="BASE_YEAR_REDUCTION",
    target_reference_std="BASE_YEAR",
    main_target_pct_std=17.5,
    unconditional_pct_std=12.8,
    conditional_total_pct_std=17.5,
    base_year_std=2014,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
    target_usability_status="USABLE",
)

# Cuba
for _cycle in ["First_NDC", "Second_NDC"]:
    apply_final_audit_override(
        r"^Cuba$", _cycle,
        audit_decision_std="EXCLUDE",
        audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
        harmonization_method_std="NONE",
    )

# Ecuador
apply_final_audit_override(
    r"^Ecuador$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)
apply_final_audit_override(
    r"^Ecuador$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_BAU_TARGET",
    audit_note_std="Audit-corrected implementation start year: 2019.",
    harmonization_method_std="DIRECT_BAU_REDUCTION",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=8.0,
    unconditional_pct_std=7.0,
    conditional_total_pct_std=8.0,
    conditional_additional_pct_std=1.0,
    implementation_start_year=2019,
    implementation_end_year=2035,
    target_year_std=2035,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)

# Egypt
apply_final_audit_override(
    r"^Egypt$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)
apply_final_audit_override(
    r"^Egypt$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="SECTORAL_TARGETS_SIMPLE_MEAN",
    audit_note_std=(
        "Audit-validated unweighted mean of the three quantified 2030 "
        "conditional sectoral targets: Electricity 37%, associated Oil & "
        "Gas 65%, and Transport 7%. Mean reduction rate = 36.33%."
    ),
    harmonization_method_std="MEAN_SECTORAL_BAU_TARGET",
    main_target_rule_std="MEAN_SECTORAL_TARGETS",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=((37.0 + 65.0 + 7.0) / 3.0),
    conditional_total_pct_std=((37.0 + 65.0 + 7.0) / 3.0),
    official_bau_value_ggco2e_std=(214740.0 + 2575.0 + 124360.0),
    implementation_start_year=2015,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
    target_usability_status="USABLE_SECTOR_MATCHED",
)

# Iraq
apply_final_audit_override(
    r"^Iraq$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)

# Kiribati — piecewise BAU trajectory with 2025 interim target.
apply_final_audit_override(
    r"^Kiribati$", "First_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_MULTI_HORIZON_BAU_TARGET",
    harmonization_method_std="PIECEWISE_BAU_TARGET",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=12.8,
    interim_target_year_std=2025,
    interim_target_pct_std=13.7,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
)
apply_final_audit_override(
    r"^Kiribati$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_MULTI_HORIZON_BAU_TARGET",
    harmonization_method_std="PIECEWISE_BAU_TARGET",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=23.8,
    unconditional_pct_std=8.0,
    conditional_total_pct_std=23.8,
    conditional_additional_pct_std=15.8,
    interim_target_year_std=2025,
    interim_target_pct_std=16.7,
    interim_unconditional_pct_std=9.5,
    interim_conditional_total_pct_std=16.7,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
)

# Mali — validated unweighted mean of sectoral BAU reduction rates.
apply_final_audit_override(
    r"^Mali$", "First_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="SECTORAL_TARGETS_HARMONISED",
    audit_note_std="National proxy = unweighted mean of 29%, 31%, 21% sectoral BAU targets.",
    harmonization_method_std="MEAN_SECTORAL_BAU_TARGET",
    main_target_rule_std="MEAN_SECTORAL_TARGETS",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=27.0,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
)
apply_final_audit_override(
    r"^Mali$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="SECTORAL_TARGETS_HARMONISED",
    audit_note_std="National proxy = unweighted mean of 31%, 25%, 39%, 31% sectoral BAU targets.",
    harmonization_method_std="MEAN_SECTORAL_BAU_TARGET",
    main_target_rule_std="MEAN_SECTORAL_TARGETS",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=31.5,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
)

# Mozambique
apply_final_audit_override(
    r"^Mozambique$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)
apply_final_audit_override(
    r"^Mozambique$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="CUMULATIVE_REDUCTION_TRAJECTORY",
    audit_note_std=(
        "40 MtCO2e cumulative reduction over 2020-2025; harmonised as a "
        "linear-ramp cumulative-abatement trajectory."
    ),
    harmonization_method_std="CUMULATIVE_REDUCTION_LINEAR_RAMP",
    target_metric_std="TRAJECTORY_TARGET",
    target_reference_std="BAU",
    special_target_amount_ggco2e_std=mtco2e_to_ggco2e(40.0),
    implementation_start_year=2020,
    implementation_end_year=2025,
    target_year_std=2025,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)

# Myanmar Second NDC — cumulative mitigation contribution over 2021-2030.
# The reported 244.52 MtCO2e (unconditional) and 414.75 MtCO2e
# (conditional total) are treated as cumulative abatements over the NDC
# implementation period, NOT as annual 2030 reductions.  The final builder
# uses the same linear-ramp cumulative-abatement convention as for Mozambique,
# preserving the reported cumulative total over 2021-2030.
apply_final_audit_override(
    r"^Myanmar$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="CUMULATIVE_REDUCTION_TRAJECTORY",
    audit_note_std=(
        "Myanmar NDC2 reports total mitigation contributions of 244.52 MtCO2e "
        "unconditional and 414.75 MtCO2e conditional over 2021-2030. These "
        "amounts are cumulative abatements, not annual 2030 reductions; they "
        "are harmonised with a linear-ramp cumulative-abatement trajectory."
    ),
    harmonization_method_std="CUMULATIVE_REDUCTION_LINEAR_RAMP",
    target_metric_std="TRAJECTORY_TARGET",
    target_reference_std="BAU",
    special_target_amount_ggco2e_std=mtco2e_to_ggco2e(414.75),
    special_unconditional_target_amount_ggco2e_std=mtco2e_to_ggco2e(244.52),
    special_conditional_target_amount_ggco2e_std=mtco2e_to_ggco2e(414.75),
    implementation_start_year=2021,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)

# Niue
apply_final_audit_override(
    r"^Niue$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)

# Nepal
for _cycle in ["First_NDC", "Second_NDC"]:
    apply_final_audit_override(
        r"^Nepal$", _cycle,
        audit_decision_std="EXCLUDE",
        audit_reason_std="NO_QUANTIFIED_NATIONAL_GHG_TARGET",
        harmonization_method_std="NONE",
    )

# Nauru
for _cycle in ["First_NDC", "Second_NDC"]:
    apply_final_audit_override(
        r"^Nauru$", _cycle,
        audit_decision_std="EXCLUDE",
        audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
        harmonization_method_std="NONE",
    )

# Papua New Guinea
apply_final_audit_override(
    r"^Papua New Guinea$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)
apply_final_audit_override(
    r"^Papua New Guinea$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_NET_LULUCF_FIXED_TARGET",
    audit_note_std=(
        "Use the quantified net LULUCF target only: 2015 net emissions "
        "of 1,716.46 GgCO2e less 10,000 GgCO2e gives -8,283.54 GgCO2e "
        "in 2030; observed emissions use FAOSTAT Forest Land only."
    ),
    harmonization_method_std="FIXED_NET_LULUCF_TARGET",
    target_metric_std="FIXED_LEVEL_TARGET",
    target_reference_std="FIXED_LEVEL",
    main_target_value_ggco2e_std=-8283.54,
    main_target_value_kind_std="EMISSIONS_LEVEL_OR_CAP",
    implementation_start_year=2021,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="SINGLE_SECTOR_GHG",
    target_usability_status="USABLE_FIXED_LEVEL_SECTOR_MATCHED",
)

# DPR Korea / North Korea
for _pat in [r"^North Korea$", r"^Democratic People's Republic of Korea$", r"^DPR Korea$"]:
    apply_final_audit_override(
        _pat, "First_NDC",
        audit_decision_std="EXCLUDE",
        audit_reason_std="NO_RELIABLE_BAU_SERIES",
        audit_note_std="No reliable official BAU trajectory for a reproducible Gap-to-Target series.",
        harmonization_method_std="NONE",
        target_usability_status="EXCLUDE_NO_RELIABLE_BAU",
    )
    apply_final_audit_override(
        _pat, "Second_NDC",
        audit_decision_std="EXCLUDE",
        audit_reason_std="NO_RELIABLE_BAU_SERIES",
        audit_note_std="A rounded BAU point inferred from amounts is not a reliable official BAU trajectory.",
        harmonization_method_std="NONE",
        target_usability_status="EXCLUDE_NO_RELIABLE_BAU",
    )

# Saudi Arabia
apply_final_audit_override(
    r"^Saudi Arabia$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)
apply_final_audit_override(
    r"^Saudi Arabia$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="ABSOLUTE_ANNUAL_ABATEMENT_TARGET",
    audit_note_std="278 MtCO2e/year avoided/reduced by 2030; not a 278 Mt emissions cap.",
    harmonization_method_std="ABSOLUTE_ANNUAL_ABATEMENT_BAU",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    special_target_amount_ggco2e_std=mtco2e_to_ggco2e(278.0),
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)

# El Salvador
apply_final_audit_override(
    r"^El Salvador$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)
apply_final_audit_override(
    r"^El Salvador$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="ABSOLUTE_ANNUAL_ABATEMENT_TARGET",
    audit_note_std="2030 annual BAU abatement: 640 ktCO2e unconditional, 819 ktCO2e total with support; AFOLU cumulative target excluded.",
    harmonization_method_std="ABSOLUTE_ANNUAL_ABATEMENT_BAU",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    special_target_amount_ggco2e_std=ktco2e_to_ggco2e(819.0),
    special_unconditional_target_amount_ggco2e_std=ktco2e_to_ggco2e(640.0),
    special_conditional_target_amount_ggco2e_std=ktco2e_to_ggco2e(819.0),
    implementation_start_year=2021,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="SINGLE_SECTOR_GHG",
)

# Suriname
for _cycle in ["First_NDC", "Second_NDC"]:
    apply_final_audit_override(
        r"^Suriname$", _cycle,
        audit_decision_std="EXCLUDE",
        audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
        harmonization_method_std="NONE",
    )

# Seychelles — official BAU values reconstructed/declared during audit.
apply_final_audit_override(
    r"^Seychelles$", "First_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_BAU_TARGET",
    harmonization_method_std="DIRECT_BAU_REDUCTION_OFFICIAL_BAU",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=29.0,
    official_bau_value_ggco2e_std=(188.0 / 0.29),
    interim_target_year_std=2025,
    interim_target_pct_std=21.4,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)
apply_final_audit_override(
    r"^Seychelles$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_BAU_TARGET",
    harmonization_method_std="DIRECT_BAU_REDUCTION_OFFICIAL_BAU",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=26.4,
    official_bau_value_ggco2e_std=1110.8,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="BROAD_NATIONAL_GHG",
)

# Syria
apply_final_audit_override(
    r"^Syrian Arab Republic$|^Syria$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)

# Timor-Leste
apply_final_audit_override(
    r"^Timor-Leste$|^East Timor$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)

# Trinidad and Tobago
apply_final_audit_override(
    r"^Trinidad and Tobago$", "First_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_MULTI_SECTOR_BAU_TARGET",
    harmonization_method_std="DIRECT_BAU_REDUCTION",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=15.0,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
)
apply_final_audit_override(
    r"^Trinidad and Tobago$", "Second_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_MULTI_SECTOR_REFERENCE_TARGET",
    harmonization_method_std="DIRECT_BAU_REDUCTION",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=15.0,
    implementation_start_year=2025,
    implementation_end_year=2035,
    target_year_std=2035,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
)

# Zambia — audited headline target: 47% reduction relative to 2010.
# Apply to both Paris cycles; conditional/unconditional components already present
# in the source remain available, while the headline target is anchored to 2010.
for _cycle in ["First_NDC", "Second_NDC"]:
    apply_final_audit_override(
        r"^Zambia$", _cycle,
        audit_decision_std="INCLUDE",
        audit_reason_std="VERIFIED_BASE_YEAR_TARGET",
        harmonization_method_std="DIRECT_BASE_YEAR_REDUCTION",
        target_metric_std="BASE_YEAR_REDUCTION",
        target_reference_std="BASE_YEAR",
        main_target_pct_std=47.0,
        base_year_std=2010,
    )

# Tuvalu — both targets are quantified, but EDGAR does not report the energy
# component required to reproduce their sectoral perimeter.  National-total
# emissions are deliberately not substituted because that would introduce
# agriculture, IPPU and waste outside the quantified energy target.
apply_final_audit_override(
    r"^Tuvalu$", "First_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="MISSING_SCOPE_COMPATIBLE_EDGAR_DATA",
    audit_note_std=(
        "The 60% reduction below 2010 for the energy sector is quantified, "
        "but EDGAR contains no usable ENERGY_ALL series for Tuvalu. A "
        "scope-consistent observed-emissions series cannot be constructed."
    ),
    harmonization_method_std="DIRECT_BASE_YEAR_REDUCTION",
    target_metric_std="BASE_YEAR_REDUCTION",
    target_reference_std="BASE_YEAR",
    main_target_pct_std=60.0,
    unconditional_pct_std=60.0,
    conditional_total_pct_std=np.nan,
    base_year_std=2010,
    implementation_start_year=2020,
    implementation_end_year=2025,
    target_year_std=2025,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
    target_usability_status="MISSING_SCOPE_COMPATIBLE_EDGAR_DATA",
)
apply_final_audit_override(
    r"^Tuvalu$", "Second_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="MISSING_SCOPE_COMPATIBLE_EDGAR_DATA",
    audit_note_std=(
        "The 60% reduction below 2010 for the energy sector is quantified, "
        "but EDGAR contains no usable ENERGY_ALL series for Tuvalu. A "
        "scope-consistent observed-emissions series cannot be constructed."
    ),
    harmonization_method_std="DIRECT_BASE_YEAR_REDUCTION",
    target_metric_std="BASE_YEAR_REDUCTION",
    target_reference_std="BASE_YEAR",
    main_target_pct_std=60.0,
    unconditional_pct_std=60.0,
    conditional_total_pct_std=np.nan,
    base_year_std=2010,
    implementation_start_year=2022,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="PARTIAL_SECTORAL_GHG",
    target_usability_status="MISSING_SCOPE_COMPATIBLE_EDGAR_DATA",
)

# Uruguay — summaries establish architecture but not a usable aggregate numerical target.
for _cycle in ["First_NDC", "Second_NDC"]:
    apply_final_audit_override(
        r"^Uruguay$", _cycle,
        audit_decision_std="EXCLUDE",
        audit_reason_std="INSUFFICIENT_AGGREGATE_TARGET_QUANTIFICATION",
        audit_note_std="Architecture identified, but no aggregate numerical target was validated in audit text.",
        harmonization_method_std="NONE",
    )

# Vanuatu
apply_final_audit_override(
    r"^Vanuatu$", "First_NDC",
    audit_decision_std="INCLUDE",
    audit_reason_std="VERIFIED_SECTORAL_BAU_TARGET",
    audit_note_std="Energy-sector BAU emissions reduction = 30% (72 GgCO2e) by 2030.",
    harmonization_method_std="DIRECT_BAU_REDUCTION",
    main_target_rule_std="SECTORAL_TARGET",
    target_metric_std="BAU_REDUCTION",
    target_reference_std="BAU",
    main_target_pct_std=30.0,
    implementation_start_year=2020,
    implementation_end_year=2030,
    target_year_std=2030,
    target_year_source="MANUAL_AUDIT_VERIFIED",
    quantified_target_coverage_std="SINGLE_SECTOR_GHG",
)
apply_final_audit_override(
    r"^Vanuatu$", "Second_NDC",
    audit_decision_std="EXCLUDE",
    audit_reason_std="NO_AGGREGATABLE_QUANTIFIED_GHG_TARGET",
    harmonization_method_std="NONE",
)

# AUTHORITATIVE FINAL RAW-NDC AUDIT
# Embedded directly in the Harmonizer: no external audit dataset is needed.
# These 20 countries are excluded from both NDC cycles because the raw-NDC
# audit found no usable quantified GHG target or an incompatible architecture.
# Egypt is deliberately absent: NDC1 is excluded above, NDC2 remains included.
FINAL_RAW_NDC_EXCLUSION_PATTERNS = {
    "ATG": r"^Antigua (and|&) Barbuda$",
    "BHR": r"^Bahrain$",
    "BOL": r"^Bolivia$",
    "CUB": r"^Cuba$",
    "TLS": r"^Timor-Leste$|^East Timor$",
    "ECU": r"^Ecuador$",
    "GUY": r"^Guyana$",
    "IRN": r"^Iran$|^Iran, Islamic Republic of$",
    "IRQ": r"^Iraq$",
    "LBY": r"^Libya$",
    "NRU": r"^Nauru$",
    "NPL": r"^Nepal$",
    "PNG": r"^Papua New Guinea$",
    "SAU": r"^Saudi Arabia$",
    "SSD": r"^South Sudan$",
    "SUR": r"^Suriname$",
    "SYR": r"^Syrian Arab Republic$|^Syria$",
    "URY": r"^Uruguay$",
    "VUT": r"^Vanuatu$",
    "YEM": r"^Yemen$",
}

if len(FINAL_RAW_NDC_EXCLUSION_PATTERNS) != 20:
    raise AssertionError("The final raw-NDC exclusion list must contain 20 countries.")
if "EGY" in FINAL_RAW_NDC_EXCLUSION_PATTERNS:
    raise AssertionError("Egypt must remain eligible through its Second NDC target.")

for _iso3_audit, _country_pattern in FINAL_RAW_NDC_EXCLUSION_PATTERNS.items():
    for _cycle in ["First_NDC", "Second_NDC"]:
        apply_final_audit_override(
            _country_pattern, _cycle,
            audit_decision_std="EXCLUDE",
            audit_reason_std="FINAL_RAW_NDC_AUDIT_NO_USABLE_OR_COMPATIBLE_TARGET",
            audit_note_std=(
                "Country excluded from the final analytical panel after the "
                "country-by-country raw-NDC audit."
            ),
            harmonization_method_std="NONE",
            target_usability_status="AUDIT_EXCLUDED",
        )

# Yemen: if an INDC-only row exists in the source, it is not part of the NDC1/NDC2
# analytical sample. No override is needed when no First/Second NDC row exists.

# Corrections validées le 28 août 2026. Les champs source restent inchangés.
def apply_architecture_audit(data):
    data = data.copy()
    countries = data[col("Country")].map(ascii_lower)
    cycles = data[col("NDCs_number")].map(clean_text)
    data["audit_schema_version"] = "20260828_DJI"
    data["scope_audit_note"] = ""
    data["scope_sector_text_std"] = ""
    data["target_availability_std"] = "REPORTED"

    def update(country, cycle, **values):
        mask = countries.eq(country) & cycles.eq(cycle)
        for key, value in values.items():
            data.loc[mask, key] = value

    update("bangladesh", "Second_NDC", conditionality_std="MIXED_CONDITIONALITY")
    update("egypt", "Second_NDC", conditionality_std="CONDITIONAL_ONLY",
           target_scope_std="MULTISECTORAL", scope_audit_note="Quantified electricity, oil/gas and transport targets only.")
    update("el salvador", "Second_NDC", conditionality_std="MIXED_CONDITIONALITY",
           target_scope_std="SECTORAL", scope_audit_note="Annual 2030 target covers energy combustion; cumulative AFOLU excluded.")
    update("guatemala", "Second_NDC", target_metric_std="BASE_YEAR_REDUCTION",
           target_reference_std="BASE_YEAR", base_year_std=2016, target_year_std=2030,
           implementation_start_year=2022, implementation_end_year=2030,
           main_target_pct_std=22.6, unconditional_pct_std=11.2, conditional_total_pct_std=22.6,
           conditional_additional_pct_std=np.nan, main_target_rule_std="CONDITIONAL_TOTAL",
           conditionality_std="MIXED_CONDITIONALITY", target_scope_std="ECONOMY_WIDE_OR_BROAD",
           quantified_target_coverage_std="BROAD_NATIONAL_GHG", harmonization_method_std="DIRECT_BASE_YEAR_REDUCTION",
           audit_decision_std="INCLUDE", audit_reason_std="USER_VALIDATED_BASE_YEAR_2016",
           audit_note_std="User audit decision: base-year 2016 and broad national coverage, overriding source BAU/Partial Sectors labels.",
           scope_audit_note="Broad national coverage: user-validated interpretation.")
    update("equatorial guinea", "First_NDC", target_scope_std="ECONOMY_WIDE_OR_BROAD",
           scope_audit_note="Audited Spanish sector list: Energy, IPPU, agriculture, forestry and waste.")
    update("bhutan", "Second_NDC", target_scope_std="ECONOMY_WIDE_OR_BROAD",
           scope_audit_note="All gases and sectors explicitly covered.")
    update("algeria", "First_NDC", target_scope_std="ECONOMY_WIDE_OR_BROAD",
           quantified_target_coverage_std="BROAD_NATIONAL_GHG",
           scope_sector_text_std="Energy; Industrial processes; Agriculture; Forestry; Waste",
           scope_audit_note="French sector list covers Energy, IPPU, AFOLU and waste.")
    update("democratic republic of the congo", "First_NDC", target_scope_std="MULTISECTORAL",
           quantified_target_coverage_std="PARTIAL_SECTORAL_GHG", audit_decision_std="INCLUDE",
           scope_sector_text_std="Energy; Agriculture; LULUCF",
           scope_audit_note="Energy, agriculture and LULUCF only; IPPU and waste explicitly excluded.")

    # Une intensité nationale sans restriction sectorielle est classée large par convention.
    intensity = data["target_metric_std"].eq("INTENSITY_TARGET")
    unspecified = data[col("Sectors")].map(ascii_lower).isin(["", "not specified", "not specified."])
    restricted = data["quantified_target_coverage_std"].isin(["SINGLE_SECTOR_GHG", "PARTIAL_SECTORAL_GHG"])
    inferred = intensity & unspecified & ~restricted
    data.loc[inferred, "target_scope_std"] = "ECONOMY_WIDE_OR_BROAD"
    data.loc[inferred, "scope_audit_note"] = "Inferred from national GDP intensity; not an explicit sector declaration or gas-coverage change."
    for country, cycle in [("china", "Second_NDC"), ("uzbekistan", "First_NDC")]:
        update(country, cycle, target_scope_std="ECONOMY_WIDE_OR_BROAD",
               scope_audit_note="User-validated national intensity convention; gas coverage unchanged.")

    status = data[col("Status")].map(ascii_lower)
    no_submission = status.eq("no document submitted")
    no_submission |= countries.isin(["cook islands", "palau"]) & cycles.eq("Second_NDC")
    no_target = status.isin(["no ghg target", "no quantified target"])
    no_target |= countries.eq("somalia") & cycles.eq("First_NDC")
    # Une cible reconstruite après audit prévaut sur un ancien statut source.
    no_target &= ~data["audit_decision_std"].eq("INCLUDE")
    data.loc[no_target, "target_availability_std"] = "NO_QUANTIFIED_GHG_TARGET"
    data.loc[no_submission, "target_availability_std"] = "NO_SUBMISSION"
    # Décision d'audit : le document de 2025 n'appartient pas au deuxième cycle étudié.
    update("djibouti", "Second_NDC", target_availability_std="NO_SUBMISSION",
           audit_decision_std="EXCLUDE", audit_reason_std="NO_SUBMISSION_FOR_ANALYTICAL_NDC2",
           audit_note_std="User-validated cycle assignment: the 2025 document is treated as a later NDC edition, not analytical NDC2. Source records are preserved.")
    unavailable = data["target_availability_std"].ne("REPORTED")
    for key in ["conditionality_std", "target_scope_std"]:
        data.loc[unavailable, key] = data.loc[unavailable, "target_availability_std"]
    data["conditional_support_flag"] = data["conditionality_std"].map({
        "UNCONDITIONAL_ONLY": 0, "CONDITIONAL_ONLY": 1,
        "MIXED_CONDITIONALITY": 1, "PARTIALLY_CONDITIONAL": 1})
    return data

df = apply_architecture_audit(df)

# Recompute dependent flags after the final audit overrides.
df["main_target_constructed_flag"] = (
    df["main_target_rule_std"].fillna("")
    .str.contains(
        "AVERAGE|MEAN|MIDPOINT|RANGE|SUM_|ADDITIONAL|AGGREGAT",
        regex=True
    )
    .astype(int)
)

df["target_range_flag"] = (
    df["main_target_rule_std"].fillna("")
    .str.contains("RANGE|AVERAGE|MEAN|MIDPOINT", regex=True)
    .astype(int)
)

df["implementation_duration_years"] = np.where(
    df["implementation_start_year"].notna()
    & df["implementation_end_year"].notna(),
    df["implementation_end_year"] - df["implementation_start_year"] + 1,
    np.nan
)

df["post2030_target_flag"] = (
    (df["target_year_std"] > 2030).fillna(False).astype(int)
)

# Re-evaluate usability using the general rules, then let the manually audited
# decision take precedence. This preserves sector coverage without forcing
# manually validated sectoral proxy targets out of the sample.
df["target_usability_status"] = df.apply(target_usability_status, axis=1)

_audit_excl = df["audit_decision_std"].eq("EXCLUDE")
_audit_incl = df["audit_decision_std"].eq("INCLUDE")

df.loc[_audit_excl, "target_usability_status"] = "AUDIT_EXCLUDED"

# Included targets that begin after the EDGAR observation window remain
# methodologically valid but not yet observable in the 2015-2024 panel.
_audit_incl_future = (
    _audit_incl
    & df["implementation_start_year"].notna()
    & (df["implementation_start_year"] > 2024)
)
df.loc[
    _audit_incl & ~_audit_incl_future,
    "target_usability_status"
] = "USABLE_AUDIT_INCLUDED"
df.loc[
    _audit_incl_future,
    "target_usability_status"
] = "NOT_YET_OBSERVABLE"

# Recompute all fields whose inputs may have changed during the final audit.
df["unconditional_share_main"] = np.where(
    df["main_target_pct_std"].notna()
    & (df["main_target_pct_std"] != 0)
    & df["unconditional_pct_std"].notna(),
    df["unconditional_pct_std"] / df["main_target_pct_std"],
    np.nan
)

df["absolute_target_available_flag"] = (
    df["main_target_value_ggco2e_std"].notna().astype(int)
)

df["quantified_ghg_target_flag"] = (
    df["main_target_pct_std"].notna()
    | df["main_target_value_ggco2e_std"].notna()
    | df["target_metric_std"].isin(
        ["FIXED_LEVEL_TARGET", "TRAJECTORY_TARGET",
         "INTENSITY_TARGET", "PER_CAPITA_TARGET",
         "BAU_PER_CAPITA_TARGET"]
    )
).astype(int)

df["strict_comparable_pct_target_flag"] = (
    df["analysis_ndc12"].eq(1)
    & df["main_target_pct_std"].notna()
    & df["target_metric_std"].isin(
        ["BAU_REDUCTION", "BASE_YEAR_REDUCTION"]
    )
    & df["quantified_target_coverage_std"].isin(
        ["NATIONAL_GHG", "BROAD_NATIONAL_GHG"]
    )
    & df["per_capita_target_flag"].eq(0)
).astype(int)

df["target_horizon_years"] = np.where(
    df["target_year_std"].notna()
    & df["submission_year_std"].notna(),
    df["target_year_std"] - df["submission_year_std"],
    np.nan
)

df["target_manual_review_flag"] = df["target_usability_status"].str.contains(
    "REVIEW|UNCLEAR|MISSING|INSUFFICIENT", regex=True, na=False
).astype(int)

# 14. COMPARAISON FIRST -> SECOND NDC

analysis = df[df["analysis_ndc12"].eq(1)].copy()

# Résumé pays-cycle
keep_cols = [
    country_col, col("NDCs_number"),
    "main_target_pct_std", "target_metric_std",
    "target_reference_std", "target_scope_std",
    "base_year_std", "target_year_std",
    "quantified_ghg_target_flag"
]

pair = analysis[keep_cols].drop_duplicates(
    [country_col, col("NDCs_number")], keep="first"
)

first = pair[pair[col("NDCs_number")].eq("First_NDC")].set_index(country_col)
second = pair[pair[col("NDCs_number")].eq("Second_NDC")].set_index(country_col)

countries = sorted(set(first.index) | set(second.index))
pair_info = {}

for country in countries:
    f = first.loc[country] if country in first.index else None
    s = second.loc[country] if country in second.index else None

    both_quant = (
        f is not None and s is not None
        and int(f["quantified_ghg_target_flag"]) == 1
        and int(s["quantified_ghg_target_flag"]) == 1
    )

    metric_changed = np.nan
    reference_changed = np.nan
    scope_changed = np.nan
    target_year_changed = np.nan
    comparable = 0
    change = np.nan
    direction = "NOT_COMPARABLE"

    if f is not None and s is not None:
        metric_changed = int(
            clean_text(f["target_metric_std"])
            != clean_text(s["target_metric_std"])
        )

        # La référence est jugée différente si architecture ou base year
        # diffèrent.
        reference_changed = int(
            clean_text(f["target_reference_std"])
            != clean_text(s["target_reference_std"])
            or (
                pd.notna(f["base_year_std"])
                and pd.notna(s["base_year_std"])
                and f["base_year_std"] != s["base_year_std"]
            )
        )

        scope_changed = int(
            clean_text(f["target_scope_std"])
            != clean_text(s["target_scope_std"])
        )

        target_year_changed = int(
            pd.notna(f["target_year_std"])
            and pd.notna(s["target_year_std"])
            and f["target_year_std"] != s["target_year_std"]
        )

        comparable = int(
            both_quant
            and metric_changed == 0
            and reference_changed == 0
            and scope_changed == 0
            and target_year_changed == 0
            and pd.notna(f["main_target_pct_std"])
            and pd.notna(s["main_target_pct_std"])
        )

        if comparable:
            change = (
                float(s["main_target_pct_std"])
                - float(f["main_target_pct_std"])
            )
            if change > 1e-9:
                direction = "STRONGER"
            elif change < -1e-9:
                direction = "WEAKER"
            else:
                direction = "UNCHANGED"

    pair_info[country] = {
        "pair_first_second_available": int(both_quant),
        "pair_target_comparable_flag": comparable,
        "target_change_first_to_second_pp": change,
        "ambition_change_direction": direction,
        "metric_changed_first_to_second": metric_changed,
        "reference_changed_first_to_second": reference_changed,
        "scope_changed_first_to_second": scope_changed,
        "target_year_changed_first_to_second": target_year_changed,
    }

for v in [
    "pair_first_second_available",
    "pair_target_comparable_flag",
    "target_change_first_to_second_pp",
    "ambition_change_direction",
    "metric_changed_first_to_second",
    "reference_changed_first_to_second",
    "scope_changed_first_to_second",
    "target_year_changed_first_to_second",
]:
    df[v] = df[country_col].map(
        lambda x: pair_info.get(clean_text(x), {}).get(v, np.nan)
    )

# 15. QUALITE DU CODAGE / NOTES D'AUDIT

def coding_audit(row):
    if row["analysis_ndc12"] != 1:
        return np.nan, np.nan

    notes = []
    quality = "HIGH"

    country = clean_text(row[country_col])
    cycle = clean_text(row[col("NDCs_number")])

    if (country, cycle) in reviewed_keys:
        quality = "HIGH_REVIEWED"
        notes.append(
            "Existing row completed/corrected from explicit NDC "
            "author-provided source documentation"
        )

    if row["quantified_ghg_target_flag"] == 0:
        notes.append("No harmonized quantitative percent target")
        quality = "LIMITED"

    if row["target_scope_std"] == "UNKNOWN":
        notes.append("Target scope could not be classified confidently")
        if quality not in {"LIMITED", "HIGH_REVIEWED"}:
            quality = "CHECK"

    if row["target_scope_std"] == "SECTORAL":
        notes.append(
            "Sectoral target; use only with compatible sector emissions "
            "or in dedicated sensitivity analysis"
        )
        if quality == "HIGH":
            quality = "MEDIUM"

    if row["target_metric_std"] in {
        "INTENSITY_TARGET", "FIXED_LEVEL_TARGET",
        "TRAJECTORY_TARGET", "MULTIPLE_TARGET_TYPES"
    }:
        notes.append(
            "Special target metric; use in robustness/sensitivity analyses"
        )
        if quality == "HIGH":
            quality = "MEDIUM"

    if row["main_target_constructed_flag"] == 1:
        notes.append(
            "Main target is harmonized/derived rather than a direct "
            "single stated percentage"
        )

    # Supprimer les doublons de notes, conserver l'ordre
    notes = list(dict.fromkeys(notes))
    return quality, "; ".join(notes) if notes else np.nan

audit = df.apply(coding_audit, axis=1)
df["coding_quality_flag"] = [x[0] for x in audit]
df["coding_note"] = [x[1] for x in audit]

# 16. NETTOYAGE COLONNES TEMPORAIRES

df = df.drop(
    columns=[
        "_cond_raw", "_reference_raw", "_sector_text", "_cap_raw"
    ],
    errors="ignore"
)

# 16B. HARMONISATION DU PERIMETRE DES EMISSIONS OBSERVEES
#
# Rule validated for the revised Gap-to-Target framework:
#
# 1. Explicit ECONOMY-WIDE / ALL-SECTORS quantified target
#       -> EDGAR national total.
#
# 2. Explicit SECTOR coverage
#       -> union/sum of the corresponding EDGAR IPCC-2006 categories.
#          Sector sets are combined as a UNION of IPCC codes, so a target
#          mentioning both "energy" and "transport" never double-counts
#          transport (transport is already part of Energy).
#
# 3. If the NDC explicitly uses NET LULUCF / sinks / removals accounting
#       -> add FAOSTAT annual net Forest-land CO2 as a transparent LULUCF
#          proxy, preserving the FAO sign (negative = net removal).
#
# 4. If target coverage is not explicit enough to build a defensible
#    sector-specific perimeter
#       -> fallback to EDGAR national total.
#
# IMPORTANT:
# - FAOSTAT Forest land is a PROXY, not the complete IPCC LULUCF sector.
# - IPCC 2006 is the primary sector classification. IPCC 1996 is processed
#   only as a consistency/documentation layer and is not added to IPCC 2006.
# - All observed-emissions values produced here are in Gg CO2e-equivalent
#   except the FAO component itself, which is CO2 in Gg and is added as the
#   explicit forest/LULUCF proxy when relevant.

SCOPE_SHEET_RULES = "NDC_Scope_Rules"
SCOPE_SHEET_OBS = "NDC_Scope_Observed"
SCOPE_SHEET_EDGAR06 = "EDGAR_Groups_IPCC2006"
SCOPE_SHEET_EDGAR96 = "EDGAR_Groups_IPCC1996"
SCOPE_SHEET_FAO = "FAO_Forest_Proxy"
SCOPE_SHEET_CODEBOOK = "Scope_Codebook"

def normalize_country_name(x):
    """ASCII/lowercase country key for conservative cross-source matching."""
    s = ascii_lower(x)
    s = re.sub(r"\([^)]*\)", " ", s)
    s = s.replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()

# NDC/FAO names that commonly differ from EDGAR country names.
# Values are EDGAR ISO3 codes. This map is deliberately explicit/auditable.
COUNTRY_ISO3_ALIASES = {
    # Current sovereign states / standard name variants
    "andorra": "AND",
    "bolivia": "BOL",
    "bolivia plurinational state of": "BOL",
    "brunei": "BRN",
    "brunei darussalam": "BRN",
    "cabo verde": "CPV",
    "cape verde": "CPV",
    "china mainland": "CHN",
    "bonaire sint eustatius and saba": "BES",
    "british virgin islands": "VGB",
    "curacao": "CUW",
    "isle of man": "IMN",
    "jersey": "JEY",
    "saint martin french part": "MAF",
    "sint maarten dutch part": "SXM",
    "svalbard and jan mayen islands": "SJM",
    "united states virgin islands": "VIR",
    "wallis and futuna islands": "WLF",
    "ascension saint helena and tristan da cunha": "SHN",
    "comoros": "COM",
    "congo": "COG",
    "congo republic of": "COG",
    "republic of the congo": "COG",
    "democratic republic of the congo": "COD",
    "congo democratic republic of the": "COD",
    "cote d ivoire": "CIV",
    "ivory coast": "CIV",
    "czech republic": "CZE",
    "czechia": "CZE",
    "eswatini": "SWZ",
    "swaziland": "SWZ",
    "holy see": "VAT",
    "vatican city": "VAT",
    "iran": "IRN",
    "iran islamic republic of": "IRN",
    "laos": "LAO",
    "lao people s democratic republic": "LAO",
    "libya": "LBY",
    "liechtenstein": "LIE",
    "micronesia": "FSM",
    "micronesia federated states of": "FSM",
    "federated states of micronesia": "FSM",
    "moldova": "MDA",
    "republic of moldova": "MDA",
    "monaco": "MCO",
    "montenegro": "MNE",
    "north korea": "PRK",
    "democratic people s republic of korea": "PRK",
    "north macedonia": "MKD",
    "macedonia": "MKD",
    "palestine": "PSE",
    "state of palestine": "PSE",
    "russia": "RUS",
    "russian federation": "RUS",
    "san marino": "SMR",
    "serbia": "SRB",
    "south korea": "KOR",
    "republic of korea": "KOR",
    "south sudan": "SSD",
    "syria": "SYR",
    "syrian arab republic": "SYR",
    "tanzania": "TZA",
    "united republic of tanzania": "TZA",
    "timor leste": "TLS",
    "east timor": "TLS",
    "turkey": "TUR",
    "turkiye": "TUR",
    "united arab emirates": "ARE",
    "uae": "ARE",
    "united kingdom": "GBR",
    "united kingdom of great britain and northern ireland": "GBR",
    "united states": "USA",
    "united states of america": "USA",
    "venezuela": "VEN",
    "venezuela bolivarian republic of": "VEN",
    "vietnam": "VNM",
    "viet nam": "VNM",
    "saint kitts and nevis": "KNA",
    "saint lucia": "LCA",
    "saint vincent and the grenadines": "VCT",
    "sao tome and principe": "STP",
    "sao tome principe": "STP",

    # Historical multi-country aggregates deliberately NOT mapped:
    # - "ethiopia pdr": historically includes Ethiopia + Eritrea in FAO data.
    # - "belgium luxembourg": historical aggregate; must not be duplicated into
    #   both BEL and LUX.
    # - "czechoslovakia": historical aggregate of later CZE + SVK.
    # - "ussr": historical aggregate of 15 successor states.
    # - "yugoslav sfr": historical aggregate of multiple successor states.
    # - "pacific islands trust territory": historical multi-territory aggregate.
    # These entries do not affect the 2015-2024 analytical window because FAO
    # supplies current-country series separately for the relevant modern states.

    # EU aggregate: retained as a documented non-country aggregate.
    # It is deliberately NOT mapped to a sovereign-state ISO3.
    # "european union 27": "EUU",
}

def _edgar_read(sheet_name):
    """
    EDGAR sheets contain 9 metadata/header rows before the actual table.
    The row at Excel line 10 is the header (pandas header=9).
    """
    if not EDGAR_FILE.exists():
        raise FileNotFoundError(
            f"EDGAR file not found: {EDGAR_FILE}. "
            "Expected it in the Row Data directory."
        )
    out = pd.read_excel(
        EDGAR_FILE,
        sheet_name=sheet_name,
        header=9,
        engine="openpyxl"
    )
    out.columns = [clean_colname(c) for c in out.columns]
    return out

def _year_cols(frame):
    return [
        c for c in frame.columns
        if re.fullmatch(r"Y_(19|20)\d{2}", str(c))
    ]

def _edgar_to_long(frame, id_cols, value_name):
    years = _year_cols(frame)
    out = frame[id_cols + years].melt(
        id_vars=id_cols,
        value_vars=years,
        var_name="year_raw",
        value_name=value_name
    )
    out["year"] = (
        out["year_raw"].astype(str).str.replace("Y_", "", regex=False)
    )
    out["year"] = pd.to_numeric(out["year"], errors="coerce").astype("Int64")
    out[value_name] = pd.to_numeric(out[value_name], errors="coerce")
    return out.drop(columns="year_raw")

# EDGAR national totals
edgar_total_raw = _edgar_read("TOTALS BY COUNTRY")
if "Substance" in edgar_total_raw.columns:
    edgar_total_raw = edgar_total_raw[
        edgar_total_raw["Substance"].astype(str).eq("GWP_100_AR5_GHG")
    ].copy()

edgar_total_long = _edgar_to_long(
    edgar_total_raw,
    ["Country_code_A3", "Name"],
    "edgar_total_ggco2e"
).rename(columns={
    "Country_code_A3": "iso3",
    "Name": "edgar_country"
})

edgar_total_long["iso3"] = edgar_total_long["iso3"].astype(str).str.strip()

# EDGAR IPCC 2006 detail
edgar06_raw = _edgar_read("IPCC 2006")
if "Substance" in edgar06_raw.columns:
    edgar06_raw = edgar06_raw[
        edgar06_raw["Substance"].astype(str).eq("GWP_100_AR5_GHG")
    ].copy()

EDGAR06_CODE_COL = "ipcc_code_2006_for_standard_report"
EDGAR06_NAME_COL = "ipcc_code_2006_for_standard_report_name"

edgar06_long = _edgar_to_long(
    edgar06_raw,
    ["Country_code_A3", "Name", EDGAR06_CODE_COL, EDGAR06_NAME_COL],
    "emissions_ggco2e"
).rename(columns={
    "Country_code_A3": "iso3",
    "Name": "edgar_country",
    EDGAR06_CODE_COL: "ipcc2006_code",
    EDGAR06_NAME_COL: "ipcc2006_name",
})

edgar06_long["iso3"] = edgar06_long["iso3"].astype(str).str.strip()
edgar06_long["ipcc2006_code"] = (
    edgar06_long["ipcc2006_code"].astype(str).str.strip()
)

def _code_prefix_mask(series, prefixes):
    """True when an IPCC code equals or begins with one of the prefixes."""
    s = series.fillna("").astype(str)
    mask = pd.Series(False, index=s.index)
    for p in prefixes:
        mask |= s.eq(p) | s.str.startswith(p + ".")
    return mask

# Non-overlapping/diagnostic groups.
EDGAR2006_GROUP_PREFIXES = {
    "energy_all": ["1"],
    "power_heat": ["1.A.1.a"],
    "energy_industries_other": ["1.A.1.bc"],
    "manufacturing_combustion": ["1.A.2"],
    "transport_all": ["1.A.3"],
    "residential_other": ["1.A.4"],
    "energy_unspecified": ["1.A.5"],
    "fugitive_all": ["1.B"],
    "ippu_all": ["2"],
    "agriculture_all": ["3"],
    "waste_all": ["4"],
    "other_edgar": ["5"],
}

def aggregate_edgar_groups(long_df, code_col, group_prefixes, prefix_func):
    base = (
        long_df[["iso3", "edgar_country", "year"]]
        .drop_duplicates()
        .sort_values(["iso3", "year"])
        .reset_index(drop=True)
    )

    for group_name, prefixes in group_prefixes.items():
        m = prefix_func(long_df[code_col], prefixes)
        agg = (
            long_df.loc[m]
            .groupby(["iso3", "edgar_country", "year"], as_index=False)["emissions_ggco2e"]
            .sum(min_count=1)
            .rename(columns={"emissions_ggco2e": f"edgar_{group_name}_ggco2e"})
        )
        base = base.merge(
            agg,
            on=["iso3", "edgar_country", "year"],
            how="left"
        )
    return base

edgar_groups_2006 = aggregate_edgar_groups(
    edgar06_long,
    "ipcc2006_code",
    EDGAR2006_GROUP_PREFIXES,
    _code_prefix_mask
).merge(
    edgar_total_long[["iso3", "year", "edgar_total_ggco2e"]],
    on=["iso3", "year"],
    how="left"
)

# EDGAR IPCC 1996 — cross-check only
edgar96_raw = _edgar_read("IPCC 1996")
if "Substance" in edgar96_raw.columns:
    edgar96_raw = edgar96_raw[
        edgar96_raw["Substance"].astype(str).eq("GWP_100_AR5_GHG")
    ].copy()

EDGAR96_CODE_COL = "ipcc_code_1996_for_standard_report"
EDGAR96_NAME_COL = "ipcc_code_1996_for_standard_report_name"

edgar96_long = _edgar_to_long(
    edgar96_raw,
    ["Country_code_A3", "Name", EDGAR96_CODE_COL, EDGAR96_NAME_COL],
    "emissions_ggco2e"
).rename(columns={
    "Country_code_A3": "iso3",
    "Name": "edgar_country",
    EDGAR96_CODE_COL: "ipcc1996_code",
    EDGAR96_NAME_COL: "ipcc1996_name",
})

edgar96_long["iso3"] = edgar96_long["iso3"].astype(str).str.strip()
edgar96_long["ipcc1996_code"] = (
    edgar96_long["ipcc1996_code"].astype(str).str.strip()
)

# Broad IPCC-1996 diagnostic groups. Not used for the main scope calculation.
EDGAR1996_GROUP_PREFIXES = {
    "energy_all": ["1"],
    "ippu_all": ["2"],
    "agriculture_all": ["4"],
    "waste_all": ["6"],
}

def _code1996_prefix_mask(series, prefixes):
    s = series.fillna("").astype(str).str.replace(".", "", regex=False)
    mask = pd.Series(False, index=s.index)
    for p in prefixes:
        p2 = str(p).replace(".", "")
        mask |= s.str.startswith(p2)
    return mask

edgar_groups_1996 = aggregate_edgar_groups(
    edgar96_long,
    "ipcc1996_code",
    EDGAR1996_GROUP_PREFIXES,
    _code1996_prefix_mask
)

# FAOSTAT Forest-land proxy
if not FAO_FOREST_FILE.exists():
    raise FileNotFoundError(
        f"FAOSTAT forest file not found: {FAO_FOREST_FILE}. "
        "Expected it in the Row Data directory."
    )

fao_raw = pd.read_excel(
    FAO_FOREST_FILE,
    sheet_name="Data",
    engine="openpyxl"
)
fao_raw.columns = [clean_colname(c) for c in fao_raw.columns]

FAO_VALUE_COL = (
    "Net emissions/removals (CO2) (Forest land) (gigagrams)"
)

if FAO_VALUE_COL not in fao_raw.columns:
    raise KeyError(
        "Expected FAOSTAT Forest-land emissions/removals column not found: "
        f"{FAO_VALUE_COL}"
    )

fao_forest_long = fao_raw[
    ["m49 code", "Country", "Item Code", "Item", "Year", FAO_VALUE_COL]
].copy()

fao_forest_long = fao_forest_long.rename(columns={
    "m49 code": "m49_code",
    "Country": "fao_country",
    "Item Code": "item_code",
    "Item": "item",
    "Year": "year",
    FAO_VALUE_COL: "fao_forest_net_co2_gg",
})

fao_forest_long["year"] = pd.to_numeric(
    fao_forest_long["year"], errors="coerce"
).astype("Int64")
fao_forest_long["fao_forest_net_co2_gg"] = pd.to_numeric(
    fao_forest_long["fao_forest_net_co2_gg"], errors="coerce"
)

# Country matching
edgar_name_to_iso3 = (
    edgar_total_long[["edgar_country", "iso3"]]
    .drop_duplicates()
    .assign(country_key=lambda x: x["edgar_country"].map(normalize_country_name))
    .drop_duplicates("country_key")
    .set_index("country_key")["iso3"]
    .to_dict()
)

def country_to_iso3(name):
    key = normalize_country_name(name)
    if key in COUNTRY_ISO3_ALIASES:
        return COUNTRY_ISO3_ALIASES[key]
    return edgar_name_to_iso3.get(key, np.nan)

fao_forest_long["iso3"] = fao_forest_long["fao_country"].map(country_to_iso3)

# Preserve unresolved FAO names for audit rather than silently dropping them.
# Keep unresolved FAO entries for audit, but distinguish NDC-relevant current
# countries from historical aggregates / dependent territories. The latter are
# not forced onto a sovereign-state ISO3 because that would contaminate series.
fao_unresolved_all = sorted(
    fao_forest_long.loc[
        fao_forest_long["iso3"].isna(), "fao_country"
    ].dropna().astype(str).unique().tolist()
)

_ndc_country_keys = set(
    df[col("Country")].dropna().astype(str).map(normalize_country_name)
)

fao_unresolved = sorted(
    name for name in fao_unresolved_all
    if normalize_country_name(name) in _ndc_country_keys
)

fao_unresolved_non_ndc = sorted(
    name for name in fao_unresolved_all
    if normalize_country_name(name) not in _ndc_country_keys
)

# NDC -> observed-emissions scope rules
def _explicit_economywide_wording(text):
    s = ascii_lower(text)
    return any(k in s for k in [
        "economy-wide", "economy wide", "entire economy",
        "all sectors", "all ipcc sectors", "all emission sectors",
        "all emitting sectors", "whole economy", "all categories",
        "100% of national emissions", "100 percent of national emissions",
        "all the economic sectors", "nationwide"
    ])

def _explicit_net_forest_lulucf(text):
    """
    Conservative flag: FAO Forest-land proxy is added only when NDC text
    explicitly couples land/forest/LULUCF coverage with net/removal/sink wording.
    """
    s = ascii_lower(text)
    land = any(k in s for k in [
        "lulucf", "lucf", "land use", "forest", "forestry", "afolu"
    ])
    net_signal = any(k in s for k in [
        "net emission", "net ghg", "net co2", "net sink",
        "sink", "sinks", "removal", "removals", "absorption",
        "absorptions", "sequestration", "carbon stock"
    ])
    return int(land and net_signal)

def _scope_sector_predicates(text):
    """
    Return a list of EDGAR IPCC-2006 code predicates implied by explicit
    sector wording. The final code set is a UNION, preventing double counting.
    """
    s = ascii_lower(text)
    tests = []

    # ENERGY: use all code 1 categories if generic Energy is explicitly named.
    generic_energy = bool(re.search(r"\benergy\b", s))
    if generic_energy:
        tests.append(("ENERGY_ALL", lambda c: c.startswith("1.")))

    # Electricity/power only when generic Energy is not already selected.
    if not generic_energy and any(k in s for k in [
        "electricity", "power generation", "power sector",
        "electricity generation"
    ]):
        tests.append((
            "POWER_HEAT",
            lambda c: c == "1.A.1.a" or c.startswith("1.A.1.a.")
        ))

    # Transport only when generic Energy is not already selected.
    if not generic_energy and any(k in s for k in [
        "transport", "transportation", "mobility",
        "road transport", "shipping", "aviation"
    ]):
        tests.append((
            "TRANSPORT",
            lambda c: c.startswith("1.A.3")
        ))

    # Residential / buildings only when generic Energy is absent.
    if not generic_energy and any(k in s for k in [
        "residential", "buildings", "commercial"
    ]):
        tests.append((
            "RESIDENTIAL_OTHER",
            lambda c: c.startswith("1.A.4")
        ))

    # Oil and gas: fugitive oil/gas + refining/other energy industries.
    if not generic_energy and any(k in s for k in [
        "oil and gas", "oil & gas", "petroleum", "natural gas",
        "flaring"
    ]):
        tests.append((
            "OIL_GAS",
            lambda c: c.startswith("1.B.2") or c.startswith("1.A.1.bc")
        ))

    # IPPU / industrial processes.
    if any(k in s for k in [
        "ippu", "industrial processes", "industrial process",
        "product use", "piup"
    ]):
        tests.append(("IPPU", lambda c: c.startswith("2.")))

    # Generic industry/manufacturing: include manufacturing combustion and
    # IPPU because NDC wording often uses "industry" for both dimensions.
    elif any(k in s for k in [
        "industry", "industrial", "manufacturing"
    ]):
        tests.append((
            "INDUSTRY_BROAD",
            lambda c: c.startswith("1.A.2") or c.startswith("2.")
        ))

    # Agriculture. AFOLU also triggers agriculture; forest component is handled
    # separately by the FAO proxy only when net/sink wording is explicit.
    if any(k in s for k in [
        "agriculture", "agricultural", "livestock", "afolu",
        "enteric fermentation", "managed soils", "rice"
    ]):
        tests.append(("AGRICULTURE", lambda c: c.startswith("3.")))

    if any(k in s for k in [
        "waste", "wastewater", "landfill", "solid waste"
    ]):
        tests.append(("WASTE", lambda c: c.startswith("4.")))

    return tests

def infer_emissions_scope_rule(row):
    """
    Determine how observed emissions must be constructed for one NDC cycle.
    Returns method, sector labels, FAO flag and an audit note.
    """
    sectors = clean_text(row.get(col("Sectors"), ""))
    extra = clean_text(row.get(col("Additional infomation"), ""))
    coverage = clean_text(row.get("quantified_target_coverage_std", ""))
    scope = clean_text(row.get("target_scope_std", ""))
    text = f"{sectors} {extra}".strip()
    audited_sectors = clean_text(row.get("scope_sector_text_std", ""))
    if audited_sectors:
        sectors = audited_sectors

    economywide = (
        coverage == "NATIONAL_GHG"
        or (_explicit_economywide_wording(text) and not audited_sectors)
    )

    fao_flag = _explicit_net_forest_lulucf(text)
    sector_tests = _scope_sector_predicates(sectors)
    sector_labels = [x[0] for x in sector_tests]

    # Explicit economy-wide wording always uses country total.
    if economywide:
        method = "EDGAR_TOTAL_PLUS_FAO_FOREST" if fao_flag else "EDGAR_TOTAL"
        note = (
            "Explicit economy-wide/all-sectors quantified target; "
            "use EDGAR national total."
        )
        if fao_flag:
            note += " Add FAOSTAT Forest-land net CO2 proxy for explicit net LULUCF/sinks."
        return method, ";".join(sector_labels), fao_flag, note

    # Explicit sector coverage: construct union of the relevant EDGAR categories.
    if sector_tests and coverage in {
        "SINGLE_SECTOR_GHG",
        "PARTIAL_SECTORAL_GHG",
        "BROAD_NATIONAL_GHG",
    }:
        method = "EDGAR_SECTORAL_PLUS_FAO_FOREST" if fao_flag else "EDGAR_SECTORAL"
        note = (
            "Explicit quantified sector coverage; observed emissions are the "
            "union of matching EDGAR IPCC-2006 categories."
        )
        if fao_flag:
            note += " Add FAOSTAT Forest-land net CO2 proxy for explicit net LULUCF/sinks."
        return method, ";".join(sector_labels), fao_flag, note

    # A broad/unclear target with explicit sector list is still sector-matched
    # when at least one defensible EDGAR sector can be identified.
    if sector_tests and scope in {"SECTORAL", "MULTISECTORAL"}:
        method = "EDGAR_SECTORAL_PLUS_FAO_FOREST" if fao_flag else "EDGAR_SECTORAL"
        note = (
            "Sector wording is explicit enough for an EDGAR IPCC-2006 scope match."
        )
        if fao_flag:
            note += " Add FAOSTAT Forest-land proxy for explicit net LULUCF/sinks."
        return method, ";".join(sector_labels), fao_flag, note

    # Validated fallback requested for ambiguous coverage.
    method = "EDGAR_TOTAL_PLUS_FAO_FOREST" if fao_flag else "EDGAR_TOTAL_FALLBACK"
    note = (
        "Quantified sector perimeter not explicit enough for a defensible "
        "sector sum; use EDGAR national total as fallback."
    )
    if fao_flag:
        note += " Add FAOSTAT Forest-land proxy because net LULUCF/sinks are explicit."
    return method, ";".join(sector_labels), fao_flag, note

scope_rows = []
for idx, row in df.loc[df["analysis_ndc12"].eq(1)].iterrows():
    method, labels, fao_flag, note = infer_emissions_scope_rule(row)
    if clean_text(row.get("scope_audit_note", "")):
        note = clean_text(row["scope_audit_note"]) + " " + note

    scope_rows.append({
        "ndc_row_index": idx,
        "Country": clean_text(row.get(col("Country"), "")),
        "iso3": country_to_iso3(row.get(col("Country"), "")),
        "NDCs_number": clean_text(row.get(col("NDCs_number"), "")),
        "quantified_target_coverage_std": clean_text(
            row.get("quantified_target_coverage_std", "")
        ),
        "target_scope_std": clean_text(row.get("target_scope_std", "")),
        "sector_text": clean_text(row.get(col("Sectors"), "")),
        "sector_mapping_text_std": clean_text(row.get("scope_sector_text_std", "")) or clean_text(row.get(col("Sectors"), "")),
        "emissions_scope_method_std": method,
        "edgar_sector_groups_std": labels,
        "fao_forest_proxy_flag": int(fao_flag),
        "emissions_scope_note_std": note,
    })

ndc_scope_rules = pd.DataFrame(scope_rows)

# COUNTRY-CYCLE OBSERVED-EMISSIONS SCOPE OVERRIDES
# Gabon First NDC (audited):
# - 50% below baseline scenario in 2025;
# - carbon stocks stored in forest biomass are explicitly excluded;
# - agriculture and industrial processes (cement) are explicitly excluded.
# Therefore the observed-emissions perimeter is EDGAR IPCC-2006 Energy (1.*)
# + Waste (4.*), without the FAOSTAT Forest-land proxy.
_gab_n1 = (
    ndc_scope_rules["Country"].map(normalize_country_name).eq("gabon")
    & ndc_scope_rules["NDCs_number"].eq("First_NDC")
)
if _gab_n1.any():
    ndc_scope_rules.loc[_gab_n1, "emissions_scope_method_std"] = "EDGAR_SECTORAL"
    ndc_scope_rules.loc[_gab_n1, "edgar_sector_groups_std"] = "ENERGY_ALL;WASTE"
    ndc_scope_rules.loc[_gab_n1, "fao_forest_proxy_flag"] = 0
    ndc_scope_rules.loc[_gab_n1, "sector_mapping_text_std"] = "Energy; Waste"
    ndc_scope_rules.loc[_gab_n1, "emissions_scope_note_std"] = (
        "MANUAL AUDIT OVERRIDE — Gabon First NDC: EDGAR IPCC-2006 Energy "
        "(1.*) + Waste (4.*). Agriculture and industrial processes are "
        "explicitly excluded. FAOSTAT Forest-land is not added because "
        "the NDC explicitly excludes carbon stocks stored in forest biomass."
    )

# ISO3 resolution and EDGAR availability are different questions. A country name
# can be harmonised to a valid ISO3 even if EDGAR has no row for that territory.
_edgar_iso_available = set(
    edgar_total_long["iso3"].dropna().astype(str).str.upper().unique()
)
# Tuvalu — audited scope: Energy sector only.  The 60% target is a reduction
# (stored as +60 in the harmonised percentage convention), relative to 2010.
_tuv = (
    ndc_scope_rules["Country"].map(normalize_country_name).eq("tuvalu")
    & ndc_scope_rules["NDCs_number"].isin(["First_NDC", "Second_NDC"])
)
if _tuv.any():
    ndc_scope_rules.loc[_tuv, "emissions_scope_method_std"] = "EDGAR_SECTORAL"
    ndc_scope_rules.loc[_tuv, "edgar_sector_groups_std"] = "ENERGY_ALL"
    ndc_scope_rules.loc[_tuv, "fao_forest_proxy_flag"] = 0
    ndc_scope_rules.loc[_tuv, "sector_mapping_text_std"] = "Energy"
    ndc_scope_rules.loc[_tuv, "emissions_scope_note_std"] = (
        "Manual audit: Tuvalu quantified target covers Energy only; "
        "use the union of EDGAR IPCC-2006 category 1 emissions. "
        "No FAOSTAT forest component is added."
    )

# Country-cycle adaptations validated during the final audit. They are coded
# directly here; no separate country-specific dataset is required.
_cok_n1 = (ndc_scope_rules["Country"].map(normalize_country_name).eq("cook islands")
           & ndc_scope_rules["NDCs_number"].eq("First_NDC"))
if _cok_n1.any():
    ndc_scope_rules.loc[_cok_n1, "emissions_scope_method_std"] = "EDGAR_SECTORAL"
    ndc_scope_rules.loc[_cok_n1, "edgar_sector_groups_std"] = "POWER_HEAT"
    ndc_scope_rules.loc[_cok_n1, "fao_forest_proxy_flag"] = 0
    ndc_scope_rules.loc[_cok_n1, "sector_mapping_text_std"] = "Electricity generation"
    ndc_scope_rules.loc[_cok_n1, "emissions_scope_note_std"] = "Cook Islands NDC1: EDGAR POWER_HEAT only."

_fsm_n1 = (ndc_scope_rules["Country"].map(normalize_country_name).eq("micronesia")
           & ndc_scope_rules["NDCs_number"].eq("First_NDC"))
if _fsm_n1.any():
    ndc_scope_rules.loc[_fsm_n1, "emissions_scope_method_std"] = "EDGAR_SECTORAL"
    ndc_scope_rules.loc[_fsm_n1, "edgar_sector_groups_std"] = "POWER_HEAT;TRANSPORT"
    ndc_scope_rules.loc[_fsm_n1, "fao_forest_proxy_flag"] = 0
    ndc_scope_rules.loc[_fsm_n1, "sector_mapping_text_std"] = "Electricity generation; Transport"
    ndc_scope_rules.loc[_fsm_n1, "emissions_scope_note_std"] = "Micronesia NDC1: EDGAR POWER_HEAT + TRANSPORT."

_fsm_n2 = (ndc_scope_rules["Country"].map(normalize_country_name).eq("micronesia")
           & ndc_scope_rules["NDCs_number"].eq("Second_NDC"))
if _fsm_n2.any():
    ndc_scope_rules.loc[_fsm_n2, "emissions_scope_method_std"] = "EDGAR_SECTORAL"
    ndc_scope_rules.loc[_fsm_n2, "edgar_sector_groups_std"] = "POWER_HEAT"
    ndc_scope_rules.loc[_fsm_n2, "fao_forest_proxy_flag"] = 0
    ndc_scope_rules.loc[_fsm_n2, "sector_mapping_text_std"] = "Electricity generation"
    ndc_scope_rules.loc[_fsm_n2, "emissions_scope_note_std"] = "Micronesia NDC2: EDGAR POWER_HEAT only."

# Bosnia and Herzegovina. NDC1 includes the national emitting sectors and
# LULUCF sinks, so the harmonised series is EDGAR total plus FAOSTAT Forest
# Land. NDC2 headline percentage targets explicitly exclude sinks and are
# therefore matched to gross EDGAR national emissions only.
_bih_n1 = (
    ndc_scope_rules["Country"].map(normalize_country_name).eq("bosnia and herzegovina")
    & ndc_scope_rules["NDCs_number"].eq("First_NDC")
)
if _bih_n1.any():
    ndc_scope_rules.loc[_bih_n1, "emissions_scope_method_std"] = "EDGAR_TOTAL_PLUS_FAO_FOREST"
    ndc_scope_rules.loc[_bih_n1, "edgar_sector_groups_std"] = "NATIONAL_TOTAL"
    ndc_scope_rules.loc[_bih_n1, "fao_forest_proxy_flag"] = 1
    ndc_scope_rules.loc[_bih_n1, "sector_mapping_text_std"] = "National emitting sectors; net forest removals"
    ndc_scope_rules.loc[_bih_n1, "emissions_scope_note_std"] = (
        "Bosnia and Herzegovina NDC1: EDGAR national total plus FAOSTAT "
        "Forest Land because the quantified accounting perimeter includes "
        "land-use change and forestry sinks."
    )

_bih_n2 = (
    ndc_scope_rules["Country"].map(normalize_country_name).eq("bosnia and herzegovina")
    & ndc_scope_rules["NDCs_number"].eq("Second_NDC")
)
if _bih_n2.any():
    ndc_scope_rules.loc[_bih_n2, "emissions_scope_method_std"] = "EDGAR_TOTAL"
    ndc_scope_rules.loc[_bih_n2, "edgar_sector_groups_std"] = "NATIONAL_TOTAL"
    ndc_scope_rules.loc[_bih_n2, "fao_forest_proxy_flag"] = 0
    ndc_scope_rules.loc[_bih_n2, "sector_mapping_text_std"] = "All emitting sectors excluding GHG sinks"
    ndc_scope_rules.loc[_bih_n2, "emissions_scope_note_std"] = (
        "Bosnia and Herzegovina NDC2: gross EDGAR national total; the "
        "headline 2030 targets explicitly exclude GHG sinks."
    )

_egy_n2 = (ndc_scope_rules["Country"].map(normalize_country_name).eq("egypt")
           & ndc_scope_rules["NDCs_number"].eq("Second_NDC"))
if _egy_n2.any():
    ndc_scope_rules.loc[_egy_n2, "emissions_scope_method_std"] = "EDGAR_SECTORAL"
    ndc_scope_rules.loc[_egy_n2, "edgar_sector_groups_std"] = "POWER_HEAT;OIL_GAS;TRANSPORT"
    ndc_scope_rules.loc[_egy_n2, "fao_forest_proxy_flag"] = 0
    ndc_scope_rules.loc[_egy_n2, "sector_mapping_text_std"] = "Electricity; Associated Oil & Gas; Transport"
    ndc_scope_rules.loc[_egy_n2, "emissions_scope_note_std"] = "Egypt NDC2: exact EDGAR union of quantified sectors."

_blz_n2 = (ndc_scope_rules["Country"].map(normalize_country_name).eq("belize")
           & ndc_scope_rules["NDCs_number"].eq("Second_NDC"))
if _blz_n2.any():
    ndc_scope_rules.loc[_blz_n2, "emissions_scope_method_std"] = "EDGAR_SECTORAL_PLUS_FAO_FOREST"
    ndc_scope_rules.loc[_blz_n2, "edgar_sector_groups_std"] = "ENERGY_ALL;AGRICULTURE;WASTE"
    ndc_scope_rules.loc[_blz_n2, "fao_forest_proxy_flag"] = 1
    ndc_scope_rules.loc[_blz_n2, "sector_mapping_text_std"] = "Energy; Agriculture; Waste; net forest removals"
    ndc_scope_rules.loc[_blz_n2, "emissions_scope_note_std"] = (
        "Belize NDC2: EDGAR Energy + Agriculture + Waste, plus FAOSTAT "
        "Forest Land because the quantified contribution explicitly includes "
        "increased AFOLU removals."
    )

_btn = (ndc_scope_rules["Country"].map(normalize_country_name).eq("bhutan")
        & ndc_scope_rules["NDCs_number"].isin(["First_NDC", "Second_NDC"]))
if _btn.any():
    ndc_scope_rules.loc[_btn, "emissions_scope_method_std"] = "EDGAR_TOTAL"
    ndc_scope_rules.loc[_btn, "edgar_sector_groups_std"] = "NATIONAL_TOTAL"
    ndc_scope_rules.loc[_btn, "fao_forest_proxy_flag"] = 0
    ndc_scope_rules.loc[_btn, "sector_mapping_text_std"] = "All emitting sectors excluding forest removals"
    ndc_scope_rules.loc[_btn, "emissions_scope_note_std"] = (
        "Bhutan: compare gross EDGAR emissions with the audited 6,300 GgCO2e "
        "neutrality cap. FAOSTAT Forest Land is not added because the sink is "
        "already the quantity defining the cap."
    )

_png_n2 = (ndc_scope_rules["Country"].map(normalize_country_name).eq("papua new guinea")
           & ndc_scope_rules["NDCs_number"].eq("Second_NDC"))
if _png_n2.any():
    ndc_scope_rules.loc[_png_n2, "emissions_scope_method_std"] = "FAO_FOREST_ONLY"
    ndc_scope_rules.loc[_png_n2, "edgar_sector_groups_std"] = ""
    ndc_scope_rules.loc[_png_n2, "fao_forest_proxy_flag"] = 1
    ndc_scope_rules.loc[_png_n2, "sector_mapping_text_std"] = "Net LULUCF / Forest land proxy"
    ndc_scope_rules.loc[_png_n2, "emissions_scope_note_std"] = "PNG NDC2: FAOSTAT Forest Land only; no EDGAR component."

ndc_scope_rules["edgar_country_available_flag"] = (
    ndc_scope_rules["iso3"].astype("string").str.upper().isin(_edgar_iso_available)
).astype(int)

ndc_scope_rules["country_match_status"] = np.select(
    [
        ndc_scope_rules["iso3"].isna(),
        ndc_scope_rules["edgar_country_available_flag"].eq(0),
    ],
    [
        "ISO3_UNRESOLVED",
        "ISO3_RESOLVED_BUT_NO_EDGAR_COUNTRY_SERIES",
    ],
    default="ISO3_RESOLVED_AND_EDGAR_AVAILABLE"
)

# Store the scope decisions back into the harmonised NDC table.
for newcol, dtype in [
    ("emissions_scope_method_std", "object"),
    ("edgar_sector_groups_std", "object"),
    ("sector_mapping_text_std", "object"),
    ("fao_forest_proxy_flag", "float"),
    ("emissions_scope_note_std", "object"),
]:
    if newcol not in df.columns:
        if dtype == "object":
            df[newcol] = pd.Series(pd.NA, index=df.index, dtype="object")
        else:
            df[newcol] = np.nan

for newcol, dtype in [
    ("edgar_country_available_flag", "float"),
    ("country_match_status", "object"),
]:
    if newcol not in df.columns:
        if dtype == "object":
            df[newcol] = pd.Series(pd.NA, index=df.index, dtype="object")
        else:
            df[newcol] = np.nan

for _, sr in ndc_scope_rules.iterrows():
    ridx = int(sr["ndc_row_index"])
    df.at[ridx, "edgar_country_available_flag"] = sr["edgar_country_available_flag"]
    df.at[ridx, "country_match_status"] = sr["country_match_status"]
    df.at[ridx, "emissions_scope_method_std"] = sr["emissions_scope_method_std"]
    df.at[ridx, "edgar_sector_groups_std"] = sr["edgar_sector_groups_std"]
    df.at[ridx, "sector_mapping_text_std"] = sr["sector_mapping_text_std"]
    df.at[ridx, "fao_forest_proxy_flag"] = sr["fao_forest_proxy_flag"]
    df.at[ridx, "emissions_scope_note_std"] = sr["emissions_scope_note_std"]

# Build country x NDC-cycle x year observed-emissions panel
fao_lookup = (
    fao_forest_long.dropna(subset=["iso3", "year"])
    .groupby(["iso3", "year"], as_index=False)["fao_forest_net_co2_gg"]
    .sum(min_count=1)
)

total_lookup = edgar_total_long[
    ["iso3", "year", "edgar_total_ggco2e"]
].copy()

# Raw sector rows keyed for fast country-year filtering.
edgar06_for_scope = edgar06_long[
    ["iso3", "year", "ipcc2006_code", "emissions_ggco2e"]
].copy()

def selected_sector_value(iso3, year, sector_text):
    tests = _scope_sector_predicates(sector_text)
    if not tests:
        return np.nan, ""

    tmp = edgar06_for_scope[
        edgar06_for_scope["iso3"].eq(str(iso3))
        & edgar06_for_scope["year"].eq(year)
    ].copy()

    if tmp.empty:
        return np.nan, ""

    # UNION across all predicates.
    selected = pd.Series(False, index=tmp.index)
    for _, predicate in tests:
        selected |= tmp["ipcc2006_code"].fillna("").astype(str).map(predicate)

    vals = pd.to_numeric(
        tmp.loc[selected, "emissions_ggco2e"], errors="coerce"
    )
    if vals.notna().sum() == 0:
        return np.nan, ""

    codes_used = sorted(
        tmp.loc[selected & tmp["emissions_ggco2e"].notna(), "ipcc2006_code"]
        .dropna().astype(str).unique().tolist()
    )
    return float(vals.sum(min_count=1)), ";".join(codes_used)

scope_obs_rows = []
available_years = sorted(
    int(y) for y in edgar_total_long["year"].dropna().unique()
)

total_dict = {
    (str(r.iso3), int(r.year)): r.edgar_total_ggco2e
    for r in total_lookup.dropna(subset=["iso3", "year"]).itertuples()
}
fao_dict = {
    (str(r.iso3), int(r.year)): r.fao_forest_net_co2_gg
    for r in fao_lookup.dropna(subset=["iso3", "year"]).itertuples()
}

for sr in ndc_scope_rules.itertuples(index=False):
    iso3 = sr.iso3
    if pd.isna(iso3):
        continue

    for year in available_years:
        total_val = total_dict.get((str(iso3), int(year)), np.nan)
        forest_val = fao_dict.get((str(iso3), int(year)), np.nan)

        method = sr.emissions_scope_method_std
        sector_val = np.nan
        codes_used = ""

        if method == "FAO_FOREST_ONLY":
            edgar_component = 0.0
        elif method.startswith("EDGAR_SECTORAL"):
            sector_val, codes_used = selected_sector_value(
                iso3, year, sr.sector_mapping_text_std
            )
            edgar_component = sector_val
        else:
            edgar_component = total_val

        fao_component = (
            forest_val if int(sr.fao_forest_proxy_flag) == 1 else 0.0
        )

        # Do not silently turn a missing required FAO value into zero.
        if int(sr.fao_forest_proxy_flag) == 1 and pd.isna(forest_val):
            observed = np.nan
            status = "MISSING_FAO_FOREST_PROXY"
        elif pd.isna(edgar_component):
            observed = np.nan
            status = "MISSING_EDGAR_COMPONENT"
        else:
            observed = float(edgar_component) + float(fao_component)
            status = "OK"

        scope_obs_rows.append({
            "Country": sr.Country,
            "iso3": iso3,
            "NDCs_number": sr.NDCs_number,
            "year": year,
            "emissions_scope_method_std": method,
            "edgar_sector_groups_std": sr.edgar_sector_groups_std,
            "sector_mapping_text_std": sr.sector_mapping_text_std,
            "edgar_ipcc2006_codes_used": codes_used,
            "fao_forest_proxy_flag": sr.fao_forest_proxy_flag,
            "edgar_component_ggco2e": edgar_component,
            "fao_forest_net_co2_gg": (
                forest_val if int(sr.fao_forest_proxy_flag) == 1 else np.nan
            ),
            "observed_emissions_scope_ggco2e": observed,
            "scope_data_status": status,
        })

ndc_scope_observed = pd.DataFrame(scope_obs_rows)
ndc_scope_observed["audit_schema_version"] = "20260828_DJI"

# Scope audit summaries
scope_method_counts = (
    ndc_scope_rules["emissions_scope_method_std"]
    .value_counts(dropna=False)
    .rename_axis("emissions_scope_method_std")
    .reset_index(name="n_ndc_cycles")
)

scope_unresolved_ndc = (
    ndc_scope_rules.loc[
        ndc_scope_rules["country_match_status"].eq("ISO3_UNRESOLVED"),
        ["Country", "NDCs_number"]
    ]
    .drop_duplicates()
    .sort_values(["Country", "NDCs_number"])
)

scope_no_edgar_country = (
    ndc_scope_rules.loc[
        ndc_scope_rules["country_match_status"].eq(
            "ISO3_RESOLVED_BUT_NO_EDGAR_COUNTRY_SERIES"
        ),
        ["Country", "iso3", "NDCs_number"]
    ]
    .drop_duplicates()
    .sort_values(["Country", "NDCs_number"])
)

scope_codebook = pd.DataFrame([
    {
        "Variable": "emissions_scope_method_std",
        "Definition": (
            "Observed-emissions perimeter rule: EDGAR_TOTAL, EDGAR_SECTORAL, "
            "optional +FAO_FOREST, or EDGAR_TOTAL_FALLBACK."
        ),
        "Methodological role": (
            "Matches observed emissions to the quantified NDC perimeter."
        ),
    },
    {
        "Variable": "edgar_sector_groups_std",
        "Definition": (
            "Human-readable sector groups inferred from explicit NDC sector wording."
        ),
        "Methodological role": (
            "Documents which EDGAR IPCC-2006 sector families are used."
        ),
    },
    {
        "Variable": "sector_mapping_text_std",
        "Definition": (
            "Audited sector text used to map the quantified NDC perimeter "
            "to EDGAR IPCC-2006 codes; equals raw sector text unless an "
            "explicit country-cycle audit override is required."
        ),
        "Methodological role": (
            "Preserves source text while documenting the exact computational "
            "sector perimeter."
        ),
    },
    {
        "Variable": "edgar_ipcc2006_codes_used",
        "Definition": (
            "Exact union of EDGAR IPCC-2006 codes entering a sector-matched "
            "country-cycle-year observation."
        ),
        "Methodological role": "Replication and audit trail; prevents double counting.",
    },
    {
        "Variable": "fao_forest_proxy_flag",
        "Definition": (
            "1 only when NDC wording explicitly couples land/forest/LULUCF "
            "with net-removal/sink accounting."
        ),
        "Methodological role": (
            "Adds FAOSTAT Forest-land net CO2 as a transparent proxy, "
            "not as complete LULUCF."
        ),
    },
    {
        "Variable": "observed_emissions_scope_ggco2e",
        "Definition": (
            "EDGAR national or sector-matched emissions plus FAOSTAT Forest-land "
            "net CO2 proxy when required."
        ),
        "Methodological role": (
            "Preferred observed-emissions series for the scope-harmonised Gap."
        ),
    },
])

# 17. CODEBOOK

codebook = [
    ("analysis_ndc12", "Binary",
     "1 for First_NDC or Second_NDC rows; 0 otherwise.",
     "Defines the priority analytical sample.",
     "From NDCs_number."),

    ("ndc_generation_num", "Integer",
     "1=First_NDC, 2=Second_NDC.",
     "NDC-cycle interactions and comparisons.",
     "From NDCs_number."),

    ("main_target_pct_std", "Percent",
     "Clean numeric version of the harmonized main target percentage.",
     "Core Gap-to-Target / ambition measure.",
     "Target in percent; fallback to Target percentage."),

    ("main_target_rule_std", "Category",
     "Normalized rule used to select/construct the main target.",
     "Auditability of harmonization decisions.",
     "Trimmed/standardized main_target_rule."),

    ("main_target_constructed_flag", "Binary",
     "1 if target is constructed using mean/midpoint/sum/component aggregation.",
     "Robustness excluding constructed targets.",
     "Derived from main_target_rule_std."),

    ("conditionality_std", "Category",
     "UNCONDITIONAL_ONLY / CONDITIONAL_ONLY / MIXED_CONDITIONALITY / "
     "PARTIALLY_CONDITIONAL / UNKNOWN.",
     "Conditionality heterogeneity tests.",
     "Normalized from Conditionality and conditionality_structure."),

    ("conditional_support_flag", "Binary",
     "1 if implementation depends wholly/partly on external support; "
     "0 if unconditional only.",
     "Mechanism/heterogeneity with climate finance.",
     "From conditionality_std."),

    ("unconditional_pct_std", "Percent",
     "Clean unconditional target percentage when reported.",
     "Domestic-effort robustness.",
     "From Unconditional_target in percent."),

    ("conditional_total_pct_std", "Percent",
     "Total target attainable under conditional/support scenario when identifiable.",
     "Maximum supported ambition.",
     "Derived conservatively from rule and components."),

    ("conditional_additional_pct_std", "Percentage points",
     "Additional conditional component beyond unconditional target when identifiable.",
     "Measures dependence on external support.",
     "Derived from stated additional component or total minus unconditional."),

    ("unconditional_share_main", "Ratio",
     "Unconditional target divided by harmonized main target.",
     "Intensity of domestic commitment.",
     "unconditional_pct_std/main_target_pct_std."),

    ("target_metric_std", "Category",
     "BAU_REDUCTION, BASE_YEAR_REDUCTION, INTENSITY_TARGET, "
     "FIXED_LEVEL_TARGET, TRAJECTORY_TARGET, MULTIPLE_TARGET_TYPES, etc.",
     "Metric-specific robustness.",
     "Normalized from GHG_target_type/Target_Type."),

    ("target_reference_std", "Category",
     "BAU / BASE_YEAR / INTENSITY_BASE / FIXED_LEVEL / TRAJECTORY / UNKNOWN.",
     "Controls reference architecture.",
     "From target metric and Reference."),

    ("quantified_target_coverage_std", "Category",
     "Coverage of the quantified target itself: NATIONAL_GHG, BROAD_NATIONAL_GHG, "
     "PARTIAL_SECTORAL_GHG, SINGLE_SECTOR_GHG, or UNCLEAR.",
     "Used for comparability with national EDGAR emissions; distinct from mitigation sectors."),
    ("target_scope_std", "Category",
     "Descriptive sector coverage: ECONOMY_WIDE_OR_BROAD / MULTISECTORAL / SECTORAL / UNKNOWN.",
     "Economy-wide-only robustness.",
     "Conservative text classification of Sectors plus target rule."),

    ("sector_count_major", "Count",
     "Number of major sector groups detected among energy, industry, "
     "agriculture/AFOLU, LULUCF/forestry, waste, transport.",
     "Breadth of NDC coverage.",
     "Keyword classification of Sectors."),

    ("sector_energy_flag", "Binary",
     "Energy/power sector covered.",
     "Sector coverage controls.", "From Sectors."),

    ("sector_industry_ipu_flag", "Binary",
     "Industry/IPPU covered.",
     "Sector coverage controls.", "From Sectors."),

    ("sector_agriculture_afolu_flag", "Binary",
     "Agriculture/AFOLU covered.",
     "Sector coverage controls.", "From Sectors."),

    ("sector_lulucf_forestry_flag", "Binary",
     "LULUCF/forestry/land-use covered.",
     "Sector coverage controls.", "From Sectors."),

    ("sector_waste_flag", "Binary",
     "Waste/wastewater covered.",
     "Sector coverage controls.", "From Sectors."),

    ("sector_transport_flag", "Binary",
     "Transport covered.",
     "Sector coverage controls.", "From Sectors."),

    ("per_capita_target_flag", "Binary",
     "1 when target is explicitly per-capita in available coded information.",
     "Exclude/identify non-standard denominators.",
     "Keyword/manual confirmation where available."),

    ("net_emissions_target_flag", "Binary",
     "1 when stored text explicitly identifies a net-emissions target.",
     "Net-vs-gross sensitivity.",
     "Conservative keyword detection; not imputed if source text absent."),

    ("target_range_flag", "Binary",
     "1 when main target derives from a range/midpoint convention.",
     "Range-target robustness.",
     "From main_target_rule_std."),

    ("implementation_start_year", "Year",
     "Parsed start year of implementation period when at least two years are present.",
     "Implementation-horizon control.",
     "Parsed from Timeframe for implementation."),

    ("implementation_end_year", "Year",
     "Parsed end year of implementation period.",
     "Implementation-horizon control.",
     "Parsed from Timeframe for implementation."),

    ("implementation_duration_years", "Years",
     "Inclusive duration of implementation period.",
     "Exposure/time-horizon control.",
     "end-start+1."),

    ("interim_target_flag", "Binary",
     "1 when an interim target percentage/year is stored.",
     "Trajectory architecture control.",
     "From interim target columns."),

    ("quantified_ghg_target_flag", "Binary",
     "1 when a quantitative percent or absolute/fixed target is available/constructed.",
     "Restrict to quantifiable commitments.",
     "From harmonized target fields and target rules."),

    ("strict_comparable_pct_target_flag", "Binary",
     "1 for broad/economy-wide BAU or base-year percentage targets, "
     "excluding sectoral and per-capita cases.",
     "Robustness check based on a homogeneous target sample.",
     "Conservative combination of scope, metric, rule and denominator."),

    ("pair_first_second_available", "Binary",
     "1 when both First and Second NDC have quantifiable targets for the country.",
     "Paired NDC analysis.",
     "Country-level pairing."),

    ("pair_target_comparable_flag", "Binary",
     "1 when First/Second targets use same metric, reference/base year, "
     "scope and target year.",
     "Valid numeric ambition-change sample.",
     "Strict pair comparability rule."),

    ("target_change_first_to_second_pp", "Percentage points",
     "Second NDC main target minus First NDC main target, only when strictly comparable.",
     "Measures ratcheting-up of ambition.",
     "Computed only if pair_target_comparable_flag=1."),

    ("ambition_change_direction", "Category",
     "STRONGER / UNCHANGED / WEAKER / NOT_COMPARABLE.",
     "Descriptive ambition transition.",
     "Sign of comparable target change."),

    ("metric_changed_first_to_second", "Binary",
     "1 if target metric changed across First/Second NDC.",
     "Explains non-comparability.",
     "Pair comparison."),

    ("reference_changed_first_to_second", "Binary",
     "1 if reference architecture/base year changed.",
     "Explains non-comparability.",
     "Pair comparison."),

    ("scope_changed_first_to_second", "Binary",
     "1 if scope category changed.",
     "Explains non-comparability.",
     "Pair comparison."),

    ("target_year_changed_first_to_second", "Binary",
     "1 if target year changed.",
     "Explains non-comparability.",
     "Pair comparison."),

    ("coding_quality_flag", "Category",
     "HIGH/HIGH_REVIEWED/MEDIUM/CHECK/LIMITED.",
     "Prioritizes manual validation.",
     "Rule-based audit indicator."),

    ("coding_note", "Text",
     "Notes on special metrics, sectoral targets, constructed targets, "
     "or explicit corrections.",
     "Audit trail for replication and country-level verification.",
     "Generated from coding diagnostics and documented manual corrections."),

    ("unconditional_target_value_ggco2e_std", "GgCO2eq",
     "Reported unconditional absolute mitigation amount or emissions level.",
     "Used when the NDC reports absolute-value components."),
    ("conditional_target_value_ggco2e_std", "GgCO2eq",
     "Reported conditional absolute mitigation amount or emissions level.",
     "Used when the NDC reports absolute-value components."),
    ("main_target_value_ggco2e_std", "GgCO2e",
     "Clean numeric absolute target/reduction value when available in the "
     "existing absolute-value field.",
     "Supports fixed-level and absolute-target cases.",
     "Parsed from Target_in_value_GgCO₂eq."),

    ("main_target_value_kind_std", "Category",
     "Identifies whether the reported absolute number is an emissions level/cap, "
     "reduction amount, or otherwise reported absolute target.",
     "Avoids mixing emission caps with abatement amounts.",
     "Derived from Cap, target rule and metric."),

    ("absolute_target_available_flag", "Binary",
     "1 if a standardized absolute target/reduction value is available.",
     "Absolute-target robustness/sample restriction.",
     "From main_target_value_ggco2e_std."),

    ("base_year_std", "Year",
     "Clean numeric base year.",
     "Base-year target reconstruction.", "From Base_year."),

    ("target_year_source", "Category",
     "Source of target year: DECLARED_TARGET_YEAR, TIMEFRAME_END_RECOVERED, or MANUAL_VERIFIED.",
     "Traceability of target-year construction."),
    ("target_year_std", "Year",
     "Clean numeric target year.",
     "Target alignment and horizon controls.", "From Target_year."),

    ("submission_year_std", "Year",
     "Calendar year of NDC submission.",
     "Controls vintage/timing of commitment.",
     "Parsed from Submission-date including Excel serial dates."),

    ("target_horizon_years", "Years",
     "Target year minus submission year.",
     "Controls ex-ante time available to meet the target.",
     "target_year_std - submission_year_std."),

    ("post2030_target_flag", "Binary",
     "1 if target year is later than 2030.",
     "Sensitivity to longer-horizon NDCs.",
     "From target_year_std."),

    ("audit_decision_std", "Category",
     "INCLUDE / EXCLUDE when the country-cycle was manually validated.",
     "Final-sample audit and construction transparency.",
     "Country-by-country NDC audit."),

    ("audit_reason_std", "Category",
     "Reason for manual inclusion/exclusion or special treatment.",
     "Documents sample construction.",
     "Country-by-country NDC audit."),

    ("audit_note_std", "Text",
     "Concise explanation of the audited target interpretation.",
     "Replication and methodological documentation.",
     "Country-by-country NDC audit."),

    ("harmonization_method_std", "Category",
     "Method used to convert the audited NDC target into an emissions target.",
     "Identifies direct, proxy, piecewise and special constructions.",
     "Country-by-country NDC audit."),

    ("official_bau_value_ggco2e_std", "GgCO2e",
     "Official/implied NDC BAU level used when explicitly verified.",
     "Preferred over PMRCPBIE for the corresponding audited target.",
     "Manual audit."),

    ("special_target_amount_ggco2e_std", "GgCO2e",
     "Main absolute annual/cumulative abatement amount for special audited architectures.",
     "Used by the final builder for special target conversions.",
     "Manual audit."),

    ("special_unconditional_target_amount_ggco2e_std", "GgCO2e",
     "Unconditional cumulative/absolute abatement amount for special audited architectures.",
     "Preserves conditionality-specific cumulative target information.",
     "Manual audit."),

    ("special_conditional_target_amount_ggco2e_std", "GgCO2e",
     "Conditional-total cumulative/absolute abatement amount for special audited architectures.",
     "Preserves conditionality-specific cumulative target information.",
     "Manual audit."),

    ("interim_target_year_std", "Year",
     "Verified intermediate target year for multi-horizon NDCs.",
     "Piecewise trajectory construction.",
     "Manual audit / NDC text."),

    ("interim_target_pct_std", "Percent",
     "Main reduction target at the verified intermediate year.",
     "Piecewise trajectory construction.",
     "Manual audit / NDC text."),

    ("interim_unconditional_pct_std", "Percent",
     "Unconditional reduction target at the intermediate year.",
     "Conditionality-specific piecewise trajectory.",
     "Manual audit / NDC text."),

    ("interim_conditional_total_pct_std", "Percent",
     "Total/conditional reduction target at the intermediate year.",
     "Conditionality-specific piecewise trajectory.",
     "Manual audit / NDC text."),

    ("emissions_scope_method_std", "Category",
     "Rule selecting observed-emissions perimeter: EDGAR national total, "
     "sector-matched IPCC-2006 sum, optional FAOSTAT Forest-land proxy, "
     "or national-total fallback.",
     "Scope-consistent Gap construction.",
     "Derived from audited quantified coverage and explicit NDC sector wording."),

    ("sector_mapping_text_std", "Text",
     "Audited sector wording used for EDGAR IPCC-2006 mapping; raw NDC sector text remains preserved.",
     "Exact computational perimeter for observed emissions.",
     "Default = raw sector text; country-cycle overrides are explicit."),

    ("edgar_sector_groups_std", "Text",
     "Sector groups used to construct sector-matched observed emissions.",
     "Audit trail / replication.",
     "Mapped to EDGAR IPCC-2006 categories; code union prevents double counting."),

    ("fao_forest_proxy_flag", "Binary",
     "1 when FAOSTAT annual net Forest-land CO2 is added because the NDC "
     "explicitly includes net LULUCF/forest sinks or removals.",
     "Proxy adjustment for NDC net-land accounting.",
     "Conservative text rule; Forest land is not presented as full LULUCF."),

    ("emissions_scope_note_std", "Text",
     "Explanation of the observed-emissions perimeter selected for the NDC cycle.",
     "Transparency and audit of emissions-scope decisions.",
     "Generated by the emissions-scope harmonisation procedure."),

]

codebook.extend([
    ("audit_schema_version", "Version", "Shared audit revision", "Pipeline consistency", "20260828_DJI"),
    ("target_availability_std", "Category", "Reported target, no quantified target, no submission, or outside window", "Separate availability from coverage and conditionality", "Source status and validated country-cycle audit"),
    ("scope_audit_note", "Text", "Evidence or convention behind corrected coverage", "Distinguish explicit and inferred coverage", "Country-cycle audit"),
    ("scope_sector_text_std", "Text", "Audited sector list when source wording needs clarification", "Override sector matching without changing raw text", "Country-cycle audit"),
])
codebook_df = pd.DataFrame(
    codebook,
    columns=[
        "Variable", "Type", "Definition",
        "Use in analysis and replication",
        "Construction / source"
    ]
)

# 17B. UNIT CONSISTENCY CHECKS
# All harmonised absolute emissions/abatement quantities passed to the builder
# must be expressed in GgCO2e. Hard-coded audit values reported in MtCO2e or
# ktCO2e are converted explicitly above using the unit helpers.
_unit_cols = [
    "main_target_value_ggco2e_std",
    "unconditional_target_value_ggco2e_std",
    "conditional_target_value_ggco2e_std",
    "official_bau_value_ggco2e_std",
    "special_target_amount_ggco2e_std",
    "special_unconditional_target_amount_ggco2e_std",
    "special_conditional_target_amount_ggco2e_std",
]
for _c in _unit_cols:
    if _c in df.columns:
        df[_c] = pd.to_numeric(df[_c], errors="coerce")

# 18. EXPORT ET MISE EN FORME

THIN_BORDER = Border(bottom=Side(style="hair", color=COLOR_GRID))

def style_header(ws, fill_color, start_col=1, end_col=None):
    end_col = end_col or ws.max_column
    fill = PatternFill("solid", fgColor=fill_color)
    for col_idx in range(start_col, end_col + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = fill
        cell.font = Font(bold=True, color=COLOR_WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 34

def style_body(ws, max_text_width=38):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    text_headers = {
        "Sectors", "Sectoral Policies", "Mitigation_policies",
        "Additional infomation", "coding_note", "Non financial support needs",
        "Financial_Mitigation", "Financial_adaptation"
    }
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        header = clean_text(ws.cell(row=1, column=col_idx).value)
        ws.column_dimensions[letter].width = max_text_width if header in text_headers else min(max(len(header)+2, 11), 24)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            cell.border = THIN_BORDER
    headers = {clean_text(ws.cell(row=1, column=j).value): j for j in range(1, ws.max_column + 1)}
    for h in text_headers:
        j = headers.get(h)
        if j:
            for col_cells in ws.iter_cols(min_col=j, max_col=j, min_row=2, max_row=ws.max_row):
                for c in col_cells:
                    c.alignment = Alignment(vertical="top", wrap_text=True)

def add_alternating_fill(ws, light_color):
    fill = PatternFill("solid", fgColor=light_color)
    for r in range(2, ws.max_row + 1):
        if r % 2 == 0:
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).fill = fill

def style_raw_sheet(ws):
    ws.sheet_properties.tabColor = COLOR_RAW
    style_header(ws, COLOR_RAW)
    style_body(ws)
    add_alternating_fill(ws, "F6F8FA")

def style_harmonized_sheet(ws):
    ws.sheet_properties.tabColor = COLOR_HARMONIZED
    if N_SOURCE_COLUMNS >= 1:
        style_header(ws, COLOR_RAW, 1, min(N_SOURCE_COLUMNS, ws.max_column))
    if ws.max_column > N_SOURCE_COLUMNS:
        style_header(ws, COLOR_HARMONIZED, N_SOURCE_COLUMNS + 1, ws.max_column)
    style_body(ws)
    if ws.max_column > N_SOURCE_COLUMNS:
        derived_fill = PatternFill("solid", fgColor=COLOR_DERIVED_LIGHT)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=N_SOURCE_COLUMNS+1, max_col=ws.max_column):
            for cell in row:
                cell.fill = derived_fill
    headers = {clean_text(ws.cell(row=1, column=j).value): j for j in range(1, ws.max_column + 1)}
    qcol = headers.get("coding_quality_flag")
    if qcol:
        letter = get_column_letter(qcol)
        rng = f"{letter}2:{letter}{ws.max_row}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${letter}2="CHECK"'], fill=PatternFill("solid", fgColor=COLOR_CHECK)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${letter}2="LIMITED"'], fill=PatternFill("solid", fgColor=COLOR_LIMITED)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${letter}2="HIGH_REVIEWED"'], fill=PatternFill("solid", fgColor=COLOR_REVIEWED)))

def style_codebook_sheet(ws):
    ws.sheet_properties.tabColor = COLOR_CODEBOOK
    style_header(ws, COLOR_CODEBOOK)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"A":34, "B":20, "C":62, "D":48, "E":52}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    fill = PatternFill("solid", fgColor=COLOR_CODEBOOK_LIGHT)
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
            if r % 2 == 0:
                cell.fill = fill

with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl", mode="w") as writer:
    raw_df.to_excel(writer, sheet_name=SHEET_RAW, index=False)
    df.to_excel(writer, sheet_name=SHEET_HARMONIZED, index=False)
    codebook_df.to_excel(writer, sheet_name=SHEET_CODEBOOK, index=False)
    wb = writer.book
    style_raw_sheet(wb[SHEET_RAW])
    style_harmonized_sheet(wb[SHEET_HARMONIZED])
    style_codebook_sheet(wb[SHEET_CODEBOOK])

# 19. EXPORT DU MODULE DE PERIMETRE DES EMISSIONS
#
# This workbook is intentionally separate from the NDC coding workbook:
# it is an intermediate, auditable bridge between raw EDGAR/FAOSTAT data
# and the final Gap-to-Target builder.
with pd.ExcelWriter(
    EMISSIONS_SCOPE_OUTPUT,
    engine="openpyxl",
    mode="w"
) as scope_writer:
    ndc_scope_rules.to_excel(
        scope_writer, sheet_name=SCOPE_SHEET_RULES, index=False
    )
    ndc_scope_observed.to_excel(
        scope_writer, sheet_name=SCOPE_SHEET_OBS, index=False
    )
    edgar_groups_2006.to_excel(
        scope_writer, sheet_name=SCOPE_SHEET_EDGAR06, index=False
    )
    edgar_groups_1996.to_excel(
        scope_writer, sheet_name=SCOPE_SHEET_EDGAR96, index=False
    )
    fao_forest_long.to_excel(
        scope_writer, sheet_name=SCOPE_SHEET_FAO, index=False
    )
    scope_codebook.to_excel(
        scope_writer, sheet_name=SCOPE_SHEET_CODEBOOK, index=False
    )

print("\n=== HARMONISATION DU PERIMETRE DES EMISSIONS ===")
print(f"Fichier créé : {EMISSIONS_SCOPE_OUTPUT.resolve()}")
print(f"Cycles NDC avec règle de périmètre : {len(ndc_scope_rules):,}")
print(f"Observations pays-cycle-année : {len(ndc_scope_observed):,}")
print("\nMéthodes de périmètre :")
print(scope_method_counts.to_string(index=False))

if len(scope_unresolved_ndc):
    print("\nPays/agrégats NDC sans ISO3 harmonisé (à auditer) :")
    print(scope_unresolved_ndc.to_string(index=False))
else:
    print("\nTous les pays souverains NDC ont un ISO3 harmonisé.")

if len(scope_no_edgar_country):
    print("\nISO3 harmonisé mais aucune série pays EDGAR disponible :")
    print(scope_no_edgar_country.to_string(index=False))

if fao_unresolved:
    print("\nNoms FAOSTAT non appariés et pertinents pour les NDC (à auditer) :")
    print("  " + ", ".join(fao_unresolved))
else:
    print("\nTous les noms FAOSTAT pertinents pour les pays NDC sont appariés.")

if fao_unresolved_non_ndc:
    print(
        "\nEntrées FAOSTAT historiques/territoriales non appariées "
        f"(sans incidence directe sur l'échantillon NDC) : {len(fao_unresolved_non_ndc)}"
    )
    print("  " + ", ".join(fao_unresolved_non_ndc[:25]))
    if len(fao_unresolved_non_ndc) > 25:
        print(f"  ... et {len(fao_unresolved_non_ndc) - 25} autres.")

print(f"Fichier créé : {OUTPUT_FILE.resolve()}")
print(f"Données brutes : {len(raw_df):,} lignes x {len(raw_df.columns):,} colonnes")
print(f"Données harmonisées : {len(df):,} lignes x {len(df.columns):,} colonnes")
print(f"Variables documentées : {len(codebook_df):,}")
print(f"Feuilles : {SHEET_RAW}, {SHEET_HARMONIZED}, {SHEET_CODEBOOK}")
